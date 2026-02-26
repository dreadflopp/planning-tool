"""MainWindow – top-level coordinator for the route planning application."""

from __future__ import annotations

import configparser
import functools
import inspect
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from PySide6.QtCore import Qt, Slot, QTimer, QRectF
from PySide6.QtGui import QAction, QColor, QFont, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen
from PySide6.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QSplitter,
    QLabel, QSpinBox, QComboBox, QPushButton, QToolButton, QMenu,
    QScrollArea, QFileDialog, QMessageBox, QInputDialog, QDialog,
    QGroupBox, QScrollBar, QSizePolicy, QFormLayout, QDialogButtonBox, QLineEdit,
)

from domain.models import (
    Route, RouteEntry, TravelSegment, EmptySpace, Visit,
    Settings, TravelMode, TravelTimeState, ExtraTimeBlock,
)
from domain.constants import (
    VISIT_WIDTH, COLUMN_SPACING,
    COLOR_HEADER_BG, COLOR_ROUTE_COLUMN_BG, COLOR_VISIT_BG, COLOR_VISIT_BORDER,
    COLOR_TRAVEL_BG, COLOR_TRAVEL_BORDER, COLOR_EMPTY_BG, COLOR_EMPTY_BORDER,
    COLOR_TRAVEL_DEFAULT_BG, COLOR_TRAVEL_DEFAULT_BORDER,
)
from services.persistence_service import PersistenceService
from services.excel_import_service import ExcelImportService
from services.map_geocoding_service import MapGeocodingService
from services.travel_time_service import TravelTimeService
from services.auto_save_manager import AutoSaveManager
from controllers.route_recalculation_engine import RouteRecalculationEngine, _t2m, _m2t, _display_time
from controllers.route_layout_engine import RouteLayoutEngine
from ui.scenes.route_scene import RouteScene
from ui.scenes.pool_scene import PoolScene
from ui.views.route_view import RouteView
from ui.views.pool_view import PoolView
from ui.flow_layout import FlowLayout
from ui.dialogs.settings_dialog import SettingsDialog
from ui.dialogs.travel_status_dialog import TravelStatusDialog
from ui.items.visit_item import VisitItem
from ui.visit_map_window import VisitMapWindow


_PROJECT_ROOT = Path(__file__).resolve().parent.parent
_SETTINGS_FILE = _PROJECT_ROOT / "settings.ini"


def _load_api_key() -> str:
    cfg = configparser.ConfigParser()
    cfg.read(str(_SETTINGS_FILE), encoding="utf-8")
    key = cfg.get("api", "key", fallback="").strip()
    if key:
        return key

    legacy_path = Path.cwd() / "settings.ini"
    try:
        same_file = legacy_path.resolve() == _SETTINGS_FILE.resolve()
    except Exception:
        same_file = False

    if not same_file and legacy_path.exists():
        legacy_cfg = configparser.ConfigParser()
        legacy_cfg.read(str(legacy_path), encoding="utf-8")
        legacy_key = legacy_cfg.get("api", "key", fallback="").strip()
        if legacy_key:
            _save_api_key(legacy_key)
            return legacy_key

    return ""


def _save_api_key(key: str):
    cfg = configparser.ConfigParser()
    cfg.read(str(_SETTINGS_FILE), encoding="utf-8")
    if "api" not in cfg:
        cfg["api"] = {}
    cfg["api"]["key"] = key
    _SETTINGS_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(_SETTINGS_FILE, "w", encoding="utf-8") as f:
        cfg.write(f)


class MainWindow(QMainWindow):
    def __init__(self, persistence: PersistenceService):
        super().__init__()
        self.setWindowTitle("Planeringsverktyg – Hemtjänst")
        self.resize(1400, 900)

        self._db = persistence
        self._settings: Settings = self._db.load_settings()
        self._apply_visit_color_palette()
        self._api_key: str = _load_api_key()
        self._default_template_address = "Angereds Torg 5, 424 65 Angered"
        self._default_templates = [
            {"name": "Kontor", "address": self._default_template_address, "duration_minutes": 10},
            {"name": "Uppstart", "address": self._default_template_address, "duration_minutes": 30},
            {"name": "Rast", "address": self._default_template_address, "duration_minutes": 40},
            {"name": "Avslut", "address": self._default_template_address, "duration_minutes": 20},
            {"type": "extra_time", "duration_minutes": self._settings.extra_time_minutes},
        ]

        # In-memory collections
        self._visits: dict[int, Visit] = {}
        self._routes: dict[int, Route] = {}

        # Services / controllers
        self._travel_svc = TravelTimeService(self._db, self._api_key, self._settings)
        self._recalc = RouteRecalculationEngine(self._db, self._travel_svc, self._settings)
        initial_scale = max(50, min(200, int(getattr(self._settings, "ui_scale_percent", 100))))
        self._settings.ui_scale_percent = initial_scale
        self._settings.font_size = max(6, min(30, int(round(12 * initial_scale / 100.0))))
        self._layout_engine = RouteLayoutEngine(self._settings.font_size)
        self._autosave = AutoSaveManager(parent=self)
        self._autosave.save_requested.connect(self._on_autosave)

        # Import / travel status helpers
        self._import_svc = ExcelImportService(self._db)
        self._travel_status: Optional[TravelStatusDialog] = None
        self._travel_debounce_timers: dict[tuple[int, int, int], QTimer] = {}
        self._failed_fallback_keys: set[tuple[str, str, str]] = set()
        self._map_window: Optional[VisitMapWindow] = None
        self._selected_visit_id: Optional[int] = None
        self._global_search_query: str = ""
        self._last_pair_alignment_conflict_count: int = 0

        # Paired visit index: visit_id → partner_visit_id
        self._pairs: dict[int, int] = {}
        self._debug_action_depth = 0
        self._debug_wrappers_installed = False
        self._bootstrap_complete = False

        self._install_debug_action_wrappers()

        # Build UI
        self._build_ui()
        self._load_data()
        self._connect_travel_signals()
        self._bootstrap_complete = True
        if self._settings.debug_mode:
            self._run_debug_integrity_scan("Init")

    def _install_debug_action_wrappers(self):
        if self._debug_wrappers_installed:
            return

        action_methods = {
            "_on_filter_changed",
            "_on_add_route",
            "_on_route_renamed",
            "_on_route_notes_changed",
            "_on_route_color_changed",
            "_on_route_shift_blocks_changed",
            "_on_route_delete",
            "_on_entry_dropped",
            "_on_entry_returned",
            "_on_extra_time_remove",
            "_on_entry_color_changed",
            "_on_entry_remove_requested",
            "_on_entry_moved",
            "_on_entry_time_edit",
            "_on_entry_duration_changed",
            "_on_empty_remove",
            "_on_travel_mode_changed",
            "_on_travel_duration_changed",
            "_on_travel_minutes_edit",
            "_on_travel_retry",
            "_on_travel_source_toggle",
            "_on_time_integrity",
            "_on_block_visibility_changed",
            "_on_pair_alignment_toggled",
            "_on_extra_time_auto_toggled",
            "_on_extra_time_minutes_changed",
            "_on_reset_all",
            "_on_default_mode_changed",
            "_on_ui_scale_percent_changed",
            "_on_import_excel",
            "_on_open_settings",
            "_on_export_state",
            "_on_import_state",
            "_on_export_excel",
            "_on_export_pdf",
            "_on_open_map",
        }

        for name in action_methods:
            fn = getattr(self, name, None)
            if not callable(fn):
                continue

            sig = inspect.signature(fn)
            positional_params = [
                p for p in sig.parameters.values()
                if p.kind in (
                    inspect.Parameter.POSITIONAL_ONLY,
                    inspect.Parameter.POSITIONAL_OR_KEYWORD,
                )
            ]
            has_var_positional = any(
                p.kind == inspect.Parameter.VAR_POSITIONAL
                for p in sig.parameters.values()
            )
            max_positional = len(positional_params)

            @functools.wraps(fn)
            def _wrapped(*args,
                         __fn=fn,
                         __name=name,
                         __has_var_positional=has_var_positional,
                         __max_positional=max_positional,
                         **kwargs):
                call_args = args if __has_var_positional else args[:__max_positional]
                self._debug_action_depth += 1
                try:
                    result = __fn(*call_args, **kwargs)
                finally:
                    self._debug_action_depth = max(0, self._debug_action_depth - 1)
                if self._debug_action_depth == 0:
                    self._on_debug_post_action(__name, call_args, kwargs)
                return result

            setattr(self, name, _wrapped)

        self._debug_wrappers_installed = True

    def _on_debug_post_action(self, action_name: str, args: tuple, kwargs: dict):
        if not self._settings.debug_mode:
            return
        if not self._bootstrap_complete:
            return
        stamp = datetime.now().strftime("%H:%M:%S")
        args_text = self._format_debug_args(args, kwargs)
        self._get_travel_status().log_debug(
            f"[{stamp}] {action_name}{args_text}"
        )
        self._run_debug_integrity_scan(action_name)

    def _format_debug_args(self, args: tuple, kwargs: dict) -> str:
        parts: list[str] = []
        for arg in args:
            text = repr(arg)
            if len(text) > 80:
                text = text[:77] + "..."
            parts.append(text)
        for k, v in kwargs.items():
            text = repr(v)
            if len(text) > 80:
                text = text[:77] + "..."
            parts.append(f"{k}={text}")
        return f"({', '.join(parts)})" if parts else "()"

    def _run_debug_integrity_scan(self, trigger: str, log_result: bool = True):
        by_route, travel_by_route, empty_by_route, issues = self._scan_time_inconsistencies()
        self._route_scene.set_inconsistent_entries(by_route)
        self._route_scene.set_inconsistent_blocks(travel_by_route, empty_by_route)
        if not log_result:
            return
        if issues:
            self._get_travel_status().log_integrity_warning(
                f"{trigger}: {len(issues)} time anomalies"
            )
            for issue in issues[:50]:
                self._get_travel_status().log_integrity_warning(issue)
            if len(issues) > 50:
                self._get_travel_status().log_integrity_warning(
                    f"... and {len(issues) - 50} more"
                )
        else:
            self._get_travel_status().log_integrity_ok(f"{trigger}: no time anomalies")

    def _scan_time_inconsistencies(self) -> tuple[
        dict[int, set[int]],
        dict[int, set[tuple[int, int]]],
        dict[int, set[tuple[int, int]]],
        list[str],
    ]:
        inconsistent_by_route: dict[int, set[int]] = {}
        inconsistent_travel_by_route: dict[int, set[tuple[int, int]]] = {}
        inconsistent_empty_by_route: dict[int, set[tuple[int, int]]] = {}
        issues: list[str] = []

        for route in sorted(self._routes.values(), key=lambda r: r.display_order):
            bad_entries: set[int] = set()
            bad_travel_pairs: set[tuple[int, int]] = set()
            bad_empty_pairs: set[tuple[int, int]] = set()
            entries = route.sorted_entries()

            for entry in entries:
                start_m = _t2m(entry.start_time)
                end_m = _t2m(entry.end_time)
                if end_m <= start_m:
                    if entry.id is not None:
                        bad_entries.add(entry.id)
                    issues.append(
                        f"{route.name}: invalid duration for '{entry.display_name}' ({_display_time(entry.start_time)}–{_display_time(entry.end_time)})"
                    )

            for idx in range(1, len(entries)):
                prev = entries[idx - 1]
                curr = entries[idx]
                seg = route.travel_segment_between(prev.id, curr.id)
                esp = route.empty_space_between(prev.id, curr.id)
                extra = route.extra_time_for_entry(curr.id)

                if seg is not None and seg.is_calculating:
                    # Ignore transient state while API travel update is in flight.
                    continue

                travel_min = max(0, seg.travel_minutes) if seg else 0
                empty_min = max(0, esp.duration_minutes) if esp else 0
                extra_min = self._settings.extra_time_minutes if extra else 0

                expected_start = _t2m(prev.end_time) + travel_min + extra_min + empty_min
                actual_start = _t2m(curr.start_time)

                if actual_start != expected_start:
                    pair_key = (prev.id, curr.id)
                    if prev.id is not None:
                        bad_entries.add(prev.id)
                    if curr.id is not None:
                        bad_entries.add(curr.id)
                    bad_travel_pairs.add(pair_key)
                    bad_empty_pairs.add(pair_key)
                    issues.append(
                        f"{route.name}: '{curr.display_name}' starts at { _display_time(curr.start_time) } but expected { _display_time(_m2t(expected_start)) }"
                    )

                if seg is not None:
                    # TravelSegment start/end are runtime display fields and may be
                    # blank right after app startup before recalculation.
                    # Only validate these fields when both are populated.
                    if seg.start_time and seg.end_time:
                        expected_seg_start = _t2m(prev.end_time)
                        expected_seg_end = expected_seg_start + travel_min
                        seg_start = _t2m(seg.start_time)
                        seg_end = _t2m(seg.end_time)
                        if seg_start != expected_seg_start or seg_end != expected_seg_end:
                            pair_key = (prev.id, curr.id)
                            if prev.id is not None:
                                bad_entries.add(prev.id)
                            if curr.id is not None:
                                bad_entries.add(curr.id)
                            bad_travel_pairs.add(pair_key)
                            issues.append(
                                f"{route.name}: travel block {prev.display_name} → {curr.display_name} is inconsistent"
                            )

                if esp is not None:
                    expected_empty_start = _t2m(prev.end_time) + travel_min + extra_min
                    expected_empty_end = expected_empty_start + empty_min
                    esp_start = _t2m(esp.start_time)
                    esp_end = _t2m(esp.end_time)
                    if esp_start != expected_empty_start or esp_end != expected_empty_end:
                        pair_key = (prev.id, curr.id)
                        if prev.id is not None:
                            bad_entries.add(prev.id)
                        if curr.id is not None:
                            bad_entries.add(curr.id)
                        bad_empty_pairs.add(pair_key)
                        issues.append(
                            f"{route.name}: gap block {prev.display_name} → {curr.display_name} is inconsistent"
                        )

            if bad_entries:
                inconsistent_by_route[route.id] = bad_entries
            if bad_travel_pairs:
                inconsistent_travel_by_route[route.id] = bad_travel_pairs
            if bad_empty_pairs:
                inconsistent_empty_by_route[route.id] = bad_empty_pairs

        return inconsistent_by_route, inconsistent_travel_by_route, inconsistent_empty_by_route, issues

    def _clear_debug_integrity_marks(self):
        self._route_scene.set_inconsistent_entries({})
        self._route_scene.set_inconsistent_blocks({}, {})

    # ------------------------------------------------------------------
    # UI construction
    # ------------------------------------------------------------------

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root_layout = QVBoxLayout(central)
        root_layout.setContentsMargins(4, 4, 4, 4)
        root_layout.setSpacing(4)

        # Toolbar
        self._build_toolbar(root_layout)

        # Filter bar – wraps to multiple rows as needed
        self._filter_bar = QWidget()
        self._filter_bar.setObjectName("filterBar")
        self._filter_bar.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred
        )
        filter_bar_vbox = QVBoxLayout(self._filter_bar)
        filter_bar_vbox.setContentsMargins(4, 2, 4, 2)
        filter_bar_vbox.setSpacing(2)

        filter_label_row = QHBoxLayout()
        self._filter_label = QLabel("Insatser:")
        filter_label_row.addWidget(self._filter_label)
        self._filter_mode_combo = QComboBox()
        self._filter_mode_combo.addItem("OR", "or")
        self._filter_mode_combo.addItem("AND", "and")
        self._filter_mode_combo.currentIndexChanged.connect(self._on_filter_changed)
        filter_label_row.addWidget(self._filter_mode_combo)

        self._filter_cb_container = QWidget()
        self._filter_cb_container.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred
        )
        self._filter_layout = FlowLayout(self._filter_cb_container, h_spacing=6, v_spacing=4)
        self._filter_layout.setContentsMargins(0, 0, 0, 0)
        filter_label_row.addWidget(self._filter_cb_container, 1)
        filter_bar_vbox.addLayout(filter_label_row)

        self._filter_buttons: dict[str, QPushButton] = {}
        root_layout.addWidget(self._filter_bar, 0)

        # Main splitter
        self._splitter = QSplitter(Qt.Orientation.Horizontal)
        self._splitter.setObjectName("mainSplitter")
        self._splitter.setChildrenCollapsible(False)
        self._splitter.setHandleWidth(10)
        self._splitter.setStyleSheet(
            """
            QSplitter#mainSplitter::handle {
                background: #A8AFB7;
            }
            QSplitter#mainSplitter::handle:hover {
                background: #6EA1D4;
            }
            """
        )
        self._splitter_start_centered = False

        # Left: routes
        self._route_scene = RouteScene(self._layout_engine)
        self._route_view = RouteView(self._route_scene)
        self._splitter.addWidget(self._route_view)

        # Right: pool
        pool_container = QWidget()
        pool_layout = QVBoxLayout(pool_container)
        pool_layout.setContentsMargins(0, 0, 0, 0)
        self._pool_scene = PoolScene(self._layout_engine)
        self._pool_view = PoolView(self._pool_scene)
        pool_layout.addWidget(self._pool_view)
        self._splitter.addWidget(pool_container)

        self._splitter.setStretchFactor(0, 3)
        self._splitter.setStretchFactor(1, 2)
        root_layout.addWidget(self._splitter, 1)  # stretch=1 so splitter takes all space

        # Sync scrollbars horizontally within each panel
        # (individual views already have their own scrollbars)

        # Wire scene signals
        self._wire_route_scene()
        self._wire_pool_scene()

    def _build_toolbar(self, parent_layout: QVBoxLayout):
        self._command_bar = QWidget(self)
        self._command_bar.setObjectName("commandBar")
        self._command_bar.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred
        )
        self._command_flow = FlowLayout(self._command_bar, h_spacing=6, v_spacing=4)
        self._command_flow.setContentsMargins(4, 2, 4, 2)
        parent_layout.addWidget(self._command_bar, 0)

        def _add_action_button(action: QAction):
            btn = QToolButton(self._command_bar)
            btn.setDefaultAction(action)
            btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextOnly)
            self._command_flow.addWidget(btn)

        # Planfil
        act_import_state = QAction("Öppna", self)
        act_import_state.triggered.connect(self._on_import_state)
        _add_action_button(act_import_state)

        act_export_state = QAction("Spara", self)
        act_export_state.triggered.connect(self._on_export_state)
        _add_action_button(act_export_state)

        act_import = QAction("Importera", self)
        act_import.triggered.connect(self._on_import_excel)
        _add_action_button(act_import)

        self._export_action_excel = QAction("Excel", self)
        self._export_action_excel.triggered.connect(self._on_export_excel)
        self._export_action_pdf = QAction("PDF", self)
        self._export_action_pdf.triggered.connect(self._on_export_pdf)
        self._export_menu = QMenu(self._command_bar)
        self._export_menu.addAction(self._export_action_excel)
        self._export_menu.addAction(self._export_action_pdf)
        self._export_button = QToolButton(self._command_bar)
        self._export_button.setText("Exportera ▾")
        self._export_button.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        self._export_button.setMenu(self._export_menu)
        self._export_button.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextOnly)
        self._command_flow.addWidget(self._export_button)

        # Planering
        act_add_route = QAction("＋ Ny rutt", self)
        act_add_route.triggered.connect(self._on_add_route)
        _add_action_button(act_add_route)

        act_integrity = QAction("Tidsintegritet", self)
        act_integrity.triggered.connect(self._on_time_integrity)
        _add_action_button(act_integrity)

        # Visning
        self._show_travel_action = QAction("Restid", self)
        self._show_travel_action.setCheckable(True)
        self._show_travel_action.setChecked(bool(self._settings.show_travel_blocks))
        self._show_travel_action.toggled.connect(self._on_block_visibility_changed)
        _add_action_button(self._show_travel_action)

        self._show_space_action = QAction("Lucka", self)
        self._show_space_action.setCheckable(True)
        self._show_space_action.setChecked(bool(self._settings.show_space_blocks))
        self._show_space_action.toggled.connect(self._on_block_visibility_changed)
        _add_action_button(self._show_space_action)

        self._show_extra_time_action = QAction("Extratid", self)
        self._show_extra_time_action.setCheckable(True)
        self._show_extra_time_action.setChecked(bool(self._settings.show_extra_time_blocks))
        self._show_extra_time_action.toggled.connect(self._on_block_visibility_changed)
        _add_action_button(self._show_extra_time_action)

        self._pair_align_action = QAction("Par-justera", self)
        self._pair_align_action.setCheckable(True)
        self._pair_align_action.setChecked(bool(self._settings.align_pair_visits))
        self._pair_align_action.toggled.connect(self._on_pair_alignment_toggled)
        _add_action_button(self._pair_align_action)

        # Sök
        self._search_toolbar_container = QWidget(self._command_bar)
        self._search_toolbar_container.setSizePolicy(
            QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed
        )
        search_row = QHBoxLayout(self._search_toolbar_container)
        search_row.setContentsMargins(0, 0, 0, 0)
        search_row.setSpacing(4)
        self._search_label = QLabel("Sök:")
        search_row.addWidget(self._search_label)
        self._global_search_edit = QLineEdit()
        self._global_search_edit.setPlaceholderText("Namn eller adress")
        self._global_search_edit.setClearButtonEnabled(False)
        self._global_search_edit.setFixedHeight(30)
        self._global_search_edit.textChanged.connect(self._on_global_search_changed)
        self._global_search_edit.setMinimumWidth(130)
        self._global_search_edit.setMaximumWidth(170)
        search_row.addWidget(self._global_search_edit)

        self._global_search_clear_btn = QPushButton("Rensa")
        self._global_search_clear_btn.clicked.connect(self._on_clear_global_search)
        search_row.addWidget(self._global_search_clear_btn)
        self._search_toolbar_container.setMaximumWidth(280)
        self._command_flow.addWidget(self._search_toolbar_container)

        # Tidsinställningar
        self._mode_toolbar_container = QWidget(self._command_bar)
        self._mode_toolbar_container.setSizePolicy(
            QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed
        )
        mode_row = QHBoxLayout(self._mode_toolbar_container)
        mode_row.setContentsMargins(0, 0, 0, 0)
        mode_row.setSpacing(4)
        self._mode_label = QLabel("Färdsätt:")
        mode_row.addWidget(self._mode_label)
        self._mode_combo = QComboBox()
        for mode, label in [(TravelMode.CAR, "Bil"),
                             (TravelMode.BIKE, "Cykel"),
                             (TravelMode.WALK, "Gå")]:
            self._mode_combo.addItem(label, mode)
        idx = self._mode_combo.findData(self._settings.default_travel_mode)
        if idx >= 0:
            self._mode_combo.setCurrentIndex(idx)
        self._mode_combo.currentIndexChanged.connect(self._on_default_mode_changed)
        mode_row.addWidget(self._mode_combo)
        self._mode_toolbar_container.setMaximumWidth(230)
        self._command_flow.addWidget(self._mode_toolbar_container)

        self._extra_time_auto_action = QAction("Auto extratid", self)
        self._extra_time_auto_action.setCheckable(True)
        self._extra_time_auto_action.setChecked(bool(self._settings.extra_time_auto_place))
        self._extra_time_auto_action.toggled.connect(self._on_extra_time_auto_toggled)
        _add_action_button(self._extra_time_auto_action)

        self._extra_time_toolbar_container = QWidget(self._command_bar)
        self._extra_time_toolbar_container.setSizePolicy(
            QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed
        )
        extra_row = QHBoxLayout(self._extra_time_toolbar_container)
        extra_row.setContentsMargins(0, 0, 0, 0)
        extra_row.setSpacing(4)
        self._extra_time_label = QLabel("Extratid:")
        extra_row.addWidget(self._extra_time_label)
        self._extra_time_spin = QSpinBox()
        self._extra_time_spin.setRange(0, 120)
        self._extra_time_spin.setSuffix(" min")
        self._extra_time_spin.setMinimumWidth(86)
        self._extra_time_spin.setMaximumWidth(96)
        self._extra_time_spin.setValue(max(0, int(self._settings.extra_time_minutes)))
        self._extra_time_spin.valueChanged.connect(self._on_extra_time_minutes_changed)
        extra_row.addWidget(self._extra_time_spin)
        self._extra_time_toolbar_container.setMaximumWidth(210)
        self._command_flow.addWidget(self._extra_time_toolbar_container)

        self._ui_scale_toolbar_container = QWidget(self._command_bar)
        self._ui_scale_toolbar_container.setSizePolicy(
            QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed
        )
        scale_row = QHBoxLayout(self._ui_scale_toolbar_container)
        scale_row.setContentsMargins(0, 0, 0, 0)
        scale_row.setSpacing(4)
        self._ui_scale_label = QLabel("Skala:")
        scale_row.addWidget(self._ui_scale_label)
        self._ui_scale_spin = QSpinBox()
        self._ui_scale_spin.setRange(50, 200)
        self._ui_scale_spin.setSingleStep(10)
        self._ui_scale_spin.setSuffix(" %")
        self._ui_scale_spin.setMinimumWidth(84)
        self._ui_scale_spin.setMaximumWidth(96)
        self._ui_scale_spin.setValue(max(50, min(200, int(self._settings.ui_scale_percent))))
        self._ui_scale_spin.valueChanged.connect(self._on_ui_scale_percent_changed)
        scale_row.addWidget(self._ui_scale_spin)
        self._ui_scale_toolbar_container.setMaximumWidth(190)
        self._command_flow.addWidget(self._ui_scale_toolbar_container)

        # System
        act_travel_log = QAction("API-logg", self)
        act_travel_log.triggered.connect(self._show_travel_status)
        _add_action_button(act_travel_log)

        act_settings = QAction("Inställningar", self)
        act_settings.triggered.connect(self._on_open_settings)
        _add_action_button(act_settings)

        act_map = QAction("Karta", self)
        act_map.triggered.connect(self._on_open_map)
        _add_action_button(act_map)

        act_reset_all = QAction("Rensa allt", self)
        act_reset_all.triggered.connect(self._on_reset_all)
        _add_action_button(act_reset_all)

    def showEvent(self, event):
        super().showEvent(event)
        if self._splitter_start_centered:
            return
        self._splitter_start_centered = True
        QTimer.singleShot(0, self._center_main_splitter)

    def _center_main_splitter(self):
        sizes = self._splitter.sizes()
        if len(sizes) < 2:
            return
        total = sizes[0] + sizes[1]
        if total <= 1:
            return
        left = max(1, total // 2)
        right = max(1, total - left)
        self._splitter.setSizes([left, right])

    def _wire_route_scene(self):
        s = self._route_scene
        s.column_reorder_requested.connect(self._on_route_column_reordered)
        s.route_renamed.connect(self._on_route_renamed)
        s.route_notes_changed.connect(self._on_route_notes_changed)
        s.route_color_changed.connect(self._on_route_color_changed)
        s.route_shift_blocks_changed.connect(self._on_route_shift_blocks_changed)
        s.route_delete_requested.connect(self._on_route_delete)
        s.entry_dropped.connect(self._on_entry_dropped)
        s.entry_moved.connect(self._on_entry_moved)
        s.entry_time_edit.connect(self._on_entry_time_edit)
        s.entry_duration_changed.connect(self._on_entry_duration_changed)
        s.entry_color_changed.connect(self._on_entry_color_changed)
        s.entry_remove_requested.connect(self._on_entry_remove_requested)
        s.empty_remove.connect(self._on_empty_remove)
        s.extra_time_remove.connect(self._on_extra_time_remove)
        s.travel_mode_changed.connect(self._on_travel_mode_changed)
        s.travel_duration_changed.connect(self._on_travel_duration_changed)
        s.travel_minutes_edit.connect(self._on_travel_minutes_edit)
        s.travel_retry.connect(self._on_travel_retry)
        s.travel_source_toggle.connect(self._on_travel_source_toggle)
        s.visit_selected.connect(self._on_route_visit_selected)
        s.entry_pair_requested.connect(self._on_route_entry_pair_requested)
        s.entry_unpair_requested.connect(self._on_route_entry_unpair_requested)
        s.pair_alignment_conflicts_changed.connect(self._on_pair_alignment_conflicts_changed)

    def _wire_pool_scene(self):
        self._pool_scene.column_order_changed.connect(self._on_pool_column_order_changed)
        self._pool_scene.entry_returned_to_pool.connect(self._on_entry_returned)
        self._pool_scene.visit_selected.connect(self._on_pool_visit_selected)
        self._pool_scene.visit_pair_requested.connect(self._on_pool_visit_pair_requested)
        self._pool_scene.visit_unpair_requested.connect(self._on_pool_visit_unpair_requested)
        self._pool_scene.template_edit_requested.connect(self._on_template_edit_requested)

    # ------------------------------------------------------------------
    # Data loading
    # ------------------------------------------------------------------

    def _load_data(self):
        visits = self._db.load_all_visits()
        self._visits = {v.id: v for v in visits}

        self._precache_map_addresses(visits)

        routes = self._db.load_all_routes(self._visits)
        self._routes = {r.id: r for r in routes}

        if self._settings.extra_time_auto_place:
            for route in routes:
                if route.extra_time_blocks:
                    continue
                entries = route.sorted_entries()
                for entry in entries[1:]:
                    route.extra_time_blocks.append(
                        ExtraTimeBlock(id=None, route_id=route.id, to_entry_id=entry.id)
                    )
                if entries:
                    self._recalc.recalculate(route)

        self._compute_pairs()
        self._build_filter_bar()
        pool_order = self._db.load_column_order("pool_street")

        # Pool: visits NOT placed in any route
        placed = self._db.get_placed_visit_ids()
        pool_visits = [v for v in visits if v.id not in placed]

        self._pool_scene.load(self._default_templates, pool_visits)
        self._pool_scene.set_column_order_map(pool_order)
        self._route_scene.load_routes(routes)
        self._pool_scene.set_extra_time_minutes(self._settings.extra_time_minutes)
        self._route_scene.set_extra_time_minutes(self._settings.extra_time_minutes)
        self._on_block_visibility_changed()
        self._route_scene.set_pair_alignment_enabled(self._settings.align_pair_visits)
        self._apply_filters_and_search()
        self._sync_map_window_visits()
        self._sync_map_selection()

    def _precache_map_addresses(self, visits: list[Visit]):
        if not self._api_key:
            return
        addresses: list[str] = []
        office_address = self._default_templates[0].get("address", "") if self._default_templates else ""
        if office_address:
            addresses.append(office_address)
        for visit in visits:
            address = (visit.full_address or visit.address or "").strip()
            if address:
                addresses.append(address)
        if not addresses:
            return
        geocode_svc = MapGeocodingService(self._db, self._api_key)
        geocode_svc.precache_addresses(addresses)

    @Slot()
    def _on_open_map(self):
        if self._map_window is None:
            office_address = self._default_templates[0].get("address", "") if self._default_templates else ""
            self._map_window = VisitMapWindow(
                self._db,
                self._api_key,
                office_address,
                log_fn=self._log_map_message,
                parent=self,
            )
            self._map_window.set_color_palette(self._settings.visit_ribbon_color_map())
            self._map_window.destroyed.connect(lambda *_: setattr(self, "_map_window", None))
            self._sync_map_window_visits()
            self._sync_map_selection()
        self._map_window.show()
        self._map_window.raise_()

    def _log_map_message(self, text: str):
        self._get_travel_status().log_debug(f"[Map] {text}")

    def _sync_map_window_visits(self):
        if self._map_window is None:
            return
        office_address = self._default_templates[0].get("address", "") if self._default_templates else ""
        self._map_window.set_office_address(office_address)
        self._map_window.set_color_palette(self._settings.visit_ribbon_color_map())
        self._map_window.set_visits(list(self._visits.values()))

    def _sync_map_selection(self):
        if self._map_window is None:
            return
        self._map_window.set_selected_visit_id(self._selected_visit_id)

    def _compute_pairs(self):
        """Find DUBBELBEMANNING 1 / DUBBELBEMANNING 2 pairs."""
        self._pairs.clear()
        cands: dict[str, list[Visit]] = {}
        for v in self._visits.values():
            ins = v.insatser.upper()
            if "DUBBELBEMANNING 1" in ins or "DUBBELBEMANNING 2" in ins:
                key = f"{v.name.strip().lower()}|{v.address.strip().lower()}"
                cands.setdefault(key, []).append(v)

        for group in cands.values():
            if len(group) < 2:
                continue
            for i in range(len(group)):
                for j in range(i + 1, len(group)):
                    a, b = group[i], group[j]
                    if self._is_recommended_pair_candidate(a, b):
                        self._pairs[a.id] = b.id
                        self._pairs[b.id] = a.id

        self._sync_visit_pair_state()

    def _sync_visit_pair_state(self):
        for visit in self._visits.values():
            visit.pair_partner_id = self._pairs.get(visit.id)

    def _is_recommended_pair_candidate(self, visit_a: Visit, visit_b: Visit) -> bool:
        """True when two visits match the auto-discovery pairing rule."""
        if not visit_a or not visit_b:
            return False
        name_a = (visit_a.name or "").strip().lower()
        name_b = (visit_b.name or "").strip().lower()
        addr_a = (visit_a.address or "").strip().lower()
        addr_b = (visit_b.address or "").strip().lower()
        if not name_a or not addr_a or name_a != name_b or addr_a != addr_b:
            return False

        a_ins = (visit_a.insatser or "").upper()
        b_ins = (visit_b.insatser or "").upper()
        a_is_1 = "DUBBELBEMANNING 1" in a_ins
        a_is_2 = "DUBBELBEMANNING 2" in a_ins
        b_is_1 = "DUBBELBEMANNING 1" in b_ins
        b_is_2 = "DUBBELBEMANNING 2" in b_ins
        valid_roles = (a_is_1 and b_is_2) or (a_is_2 and b_is_1)
        if not valid_roles:
            return False

        return abs(_t2m(visit_a.default_start) - _t2m(visit_b.default_start)) <= 5

    def _build_filter_bar(self):
        # Clear existing buttons (hide first to avoid flash)
        for btn in self._filter_buttons.values():
            btn.hide()
            btn.deleteLater()
        # Purge dead layout items
        self._filter_layout._items.clear()
        self._filter_buttons.clear()

        # Collect unique Insatser tokens
        tags: set[str] = set()
        for v in self._visits.values():
            for tag in v.insatser.split(","):
                t = tag.strip()
                if t:
                    tags.add(t)

        for tag in sorted(tags):
            btn = QPushButton(tag, self._filter_cb_container)
            btn.setCheckable(True)
            btn.setChecked(False)
            btn.toggled.connect(self._on_filter_changed)
            self._filter_buttons[tag] = btn
            self._filter_layout.addWidget(btn)

        self._filter_layout.invalidate()
        self._filter_cb_container.updateGeometry()

    # ------------------------------------------------------------------
    # Filter
    # ------------------------------------------------------------------

    @Slot()
    def _on_filter_changed(self):
        self._apply_filters_and_search()

    def _apply_filters_and_search(self):
        active = {tag for tag, btn in self._filter_buttons.items() if btn.isChecked()}
        mode = self._filter_mode_combo.currentData() or "or"
        self._route_scene.apply_filter(active, mode)
        self._pool_scene.apply_filter(active, mode)
        self._route_scene.apply_search(self._global_search_query)
        self._pool_scene.apply_search(self._global_search_query)

    @Slot(str)
    def _on_global_search_changed(self, text: str):
        self._global_search_query = (text or "").strip()
        self._apply_filters_and_search()

    @Slot()
    def _on_clear_global_search(self):
        self._global_search_edit.clear()

    @Slot(int, int)
    def _on_route_column_reordered(self, route_id: int, direction: int):
        for route in sorted(self._routes.values(), key=lambda r: r.display_order):
            self._db.save_route(route)

    @Slot()
    def _on_pool_column_order_changed(self):
        for street, order in self._pool_scene.column_order_pairs():
            self._db.save_column_order("pool_street", street, order)

    # ------------------------------------------------------------------
    # Route management
    # ------------------------------------------------------------------

    @Slot()
    def _on_add_route(self):
        name = self._next_default_route_name()
        order = max((r.display_order for r in self._routes.values()), default=-1) + 1
        route = self._db.create_route(name, order)
        self._routes[route.id] = route
        self._route_scene.add_route(route)

    def _next_default_route_name(self) -> str:
        used = {r.name for r in self._routes.values()}
        n = 1
        while True:
            candidate = f"Rutt {n}"
            if candidate not in used:
                return candidate
            n += 1

    @Slot(int, str)
    def _on_route_renamed(self, route_id: int, name: str):
        route = self._routes.get(route_id)
        if route:
            route.name = name
            self._db.save_route(route)

    @Slot(int, str)
    def _on_route_notes_changed(self, route_id: int, notes: str):
        route = self._routes.get(route_id)
        if route:
            route.notes = notes
            self._db.save_route(route)

    @Slot(int, object)
    def _on_route_color_changed(self, route_id: int, color):
        route = self._routes.get(route_id)
        if not route:
            return
        route.route_color = color or None
        for entry in route.entries:
            entry.route_color = route.route_color
        self._db.save_route(route)
        self._route_scene.rebuild_route(route, animate=False)
        self._autosave.mark_dirty(route_id)

    @Slot(int, int)
    def _on_route_shift_blocks_changed(self, route_id: int, blocks: int):
        route = self._routes.get(route_id)
        if not route:
            return
        route.vertical_shift_blocks = max(0, int(blocks))
        self._autosave.mark_dirty(route_id)

    @Slot(int)
    def _on_route_delete(self, route_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        reply = QMessageBox.question(
            self, "Ta bort rutt",
            f"Ta bort rutten '{route.name}'? Besöken återförs till listan.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        # Return all visits to pool
        for entry in route.entries:
            if entry.visit_id and entry.visit_id in self._visits:
                v = self._visits[entry.visit_id]
                self._pool_scene.add_visit(v)

        self._db.delete_route(route_id)
        del self._routes[route_id]
        self._route_scene.remove_route(route_id)

    # ------------------------------------------------------------------
    # Entry drag-and-drop
    # ------------------------------------------------------------------

    @Slot(int, object)
    def _on_entry_dropped(self, route_id: int, data: dict):
        route = self._routes.get(route_id)
        if not route:
            return

        dtype = data.get("type")
        insert_index = data.get("insert_index")
        dropped_visit_id: Optional[int] = None

        if dtype == "pool_visit":
            visit_id = data["visit_id"]
            if not self._can_place_paired_visit_on_route(route, visit_id):
                QMessageBox.warning(
                    self,
                    "Placering blockerad",
                    "Parade besök kan inte placeras på samma rutt.",
                )
                return
            visit = self._visits.get(visit_id)
            if not visit:
                return
            dur = max(1, _t2m(visit.default_end) - _t2m(visit.default_start))
            entry = RouteEntry(
                id=None, route_id=route_id, visit_id=visit_id,
                position=0,
                start_time="07:00",
                end_time=_m2t(_t2m("07:00") + dur),
                visit=visit,
            )
            self._recalc.add_entry_to_route(
                route, entry,
                self._settings.default_travel_mode,
                insert_index=insert_index,
            )
            if self._settings.extra_time_auto_place:
                self._ensure_extra_time_for_entry(route, entry.id)
            self._pool_scene.remove_visit(visit_id)
            self._request_travel_for_new_entry(route, entry)
            self._route_scene.rebuild_route(route)
            dropped_visit_id = entry.visit_id
            if insert_index is not None:
                self._route_scene.pop_visit(route_id, visit_index=int(insert_index))

        elif dtype == "office":
            duration = max(1, int(data.get("duration_minutes", 10)))
            entry = RouteEntry(
                id=None, route_id=route_id, visit_id=None,
                position=0,
                start_time="07:00", end_time=_m2t(_t2m("07:00") + duration),
                is_office_instance=True,
                office_name=data.get("name", "Kontor"),
                office_address=data.get("address", ""),
            )
            self._recalc.add_entry_to_route(
                route, entry,
                self._settings.default_travel_mode,
                insert_index=insert_index,
            )
            self._request_travel_for_new_entry(route, entry)
            self._route_scene.rebuild_route(route)
            if insert_index is not None:
                self._route_scene.pop_visit(route_id, visit_index=int(insert_index))

        elif dtype == "extra_time":
            if insert_index is None:
                insert_index = len(route.sorted_entries())
            entries = route.sorted_entries()
            if insert_index <= 0 or insert_index >= len(entries):
                return
            target = entries[insert_index]
            self._ensure_extra_time_for_entry(route, target.id)
            self._recalc.recalculate(route)
            self._route_scene.rebuild_route(route)
            self._autosave.mark_dirty(route.id)
            return

        elif dtype == "route_entry":
            # Move entry from one route to another (or reorder within same)
            entry_id = data["entry_id"]
            src_route_id = data["source_route_id"]
            if src_route_id == route_id:
                entries = route.sorted_entries()
                current_idx = next((i for i, e in enumerate(entries) if e.id == entry_id), None)
                if current_idx is None or insert_index is None:
                    return
                # insert_index is computed in the pre-removal coordinate space [0..len(entries)]
                insert_index = max(0, min(insert_index, len(entries)))

                # When moving forward, removing the source shifts later indices left by one.
                adjusted_index = insert_index - 1 if insert_index > current_idx else insert_index
                adjusted_index = max(0, min(adjusted_index, len(entries) - 1))

                if adjusted_index == current_idx:
                    return
                entry = entries.pop(current_idx)
                # After pop, list length is len(entries)-1; insert accepts index up to new len.
                adjusted_index = max(0, min(adjusted_index, len(entries)))
                entries.insert(adjusted_index, entry)
                for i, e in enumerate(entries):
                    e.position = i
                route.entries = entries
                self._recalc.recalculate(route)
                self._route_scene.rebuild_route(route)
                self._route_scene.pop_visit(route_id, visit_index=adjusted_index)
                dropped_visit_id = entry.visit_id
                if dropped_visit_id:
                    changed_route_ids = self._synchronize_all_pairs(max_passes=2)
                    for changed_route_id in changed_route_ids:
                        changed_route = self._routes.get(changed_route_id)
                        if changed_route:
                            self._route_scene.rebuild_route(changed_route)
                            self._autosave.mark_dirty(changed_route_id)
                self._autosave.mark_dirty(route.id)
                return
            src_route = self._routes.get(src_route_id)
            if not src_route:
                return
            entry = next((e for e in src_route.entries if e.id == entry_id), None)
            if not entry:
                return
            if entry.visit_id and not self._can_place_paired_visit_on_route(route, entry.visit_id):
                QMessageBox.warning(
                    self,
                    "Placering blockerad",
                    "Parade besök kan inte placeras på samma rutt.",
                )
                return
            # Remove from source
            self._recalc.remove_entry_from_route(src_route, entry, replace_with_empty=False)
            # Reset position and add to target
            entry.route_id = route_id
            entry.id = None
            self._recalc.add_entry_to_route(
                route, entry,
                self._settings.default_travel_mode,
                insert_index=insert_index,
            )
            self._request_travel_for_new_entry(route, entry)
            if self._settings.extra_time_auto_place and entry.visit_id and not entry.is_office_instance:
                self._ensure_extra_time_for_entry(route, entry.id)
            self._route_scene.rebuild_route(src_route)
            self._route_scene.rebuild_route(route)
            dropped_visit_id = entry.visit_id
            if insert_index is not None:
                self._route_scene.pop_visit(route_id, visit_index=int(insert_index))

        if dropped_visit_id:
            changed_route_ids = self._synchronize_all_pairs(max_passes=2)
            for changed_route_id in changed_route_ids:
                changed_route = self._routes.get(changed_route_id)
                if changed_route:
                    self._route_scene.rebuild_route(changed_route)
                    self._autosave.mark_dirty(changed_route_id)

        self._autosave.mark_dirty(route.id)

    @Slot(int)
    def _on_entry_returned(self, entry_id: int):
        """Visit dragged back to pool."""
        for route in self._routes.values():
            entry = next((e for e in route.entries if e.id == entry_id), None)
            if entry:
                if entry.visit_id and entry.visit_id in self._visits:
                    self._remove_entry_to_pool(route, entry)
                return

    def _remove_entry_to_pool(self, route: Route, entry: RouteEntry):
        if not entry.visit_id or entry.visit_id not in self._visits:
            return
        had_single_entry = len(route.entries) <= 1
        v = self._visits[entry.visit_id]
        v_restored = Visit(
            id=v.id, object_id=v.object_id, name=v.name,
            address=v.address, street=v.street,
            default_start=v.default_start, default_end=v.default_end,
            insatser=v.insatser, color=v.color, raw_data=v.raw_data,
            full_address=v.full_address,
            pair_partner_id=self._pairs.get(v.id),
        )
        self._recalc.remove_entry_from_route(route, entry, replace_with_empty=True)
        if had_single_entry:
            self._rebuild_all_views()
            self._autosave.mark_dirty(route.id)
            return
        self._pool_scene.add_visit(v_restored)
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route.id)

    def _ensure_extra_time_for_entry(self, route: Route, to_entry_id: Optional[int]):
        if not to_entry_id:
            return
        entries = route.sorted_entries()
        idx = next((i for i, e in enumerate(entries) if e.id == to_entry_id), None)
        if idx is None or idx == 0:
            return
        existing = route.extra_time_for_entry(to_entry_id)
        if existing:
            return
        route.extra_time_blocks.append(
            ExtraTimeBlock(id=None, route_id=route.id, to_entry_id=to_entry_id)
        )

    def _remove_extra_time_for_entry(self, route: Route, to_entry_id: int):
        block = route.extra_time_for_entry(to_entry_id)
        if not block:
            return
        route.extra_time_blocks = [
            b for b in route.extra_time_blocks if b.to_entry_id != to_entry_id
        ]

    @Slot(int, int)
    def _on_extra_time_remove(self, route_id: int, to_entry_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        self._remove_extra_time_for_entry(route, to_entry_id)
        self._recalc.recalculate(route)
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route.id)
        self._apply_pair_synchronization_and_refresh(max_passes=6)

    @Slot(int, int, object)
    def _on_entry_color_changed(self, route_id: int, entry_id: int, color):
        route = self._routes.get(route_id)
        if not route:
            return
        entry = next((e for e in route.entries if e.id == entry_id), None)
        if not entry:
            return
        if entry.is_office_instance:
            entry.office_color = color or "black"
            self._db.update_route_entry(entry)
        else:
            if not entry.visit:
                return
            entry.visit.color = color
            self._db.upsert_visit(entry.visit)
        self._route_scene.rebuild_route(route)
        self._sync_map_window_visits()

    @Slot(int, int)
    def _on_entry_remove_requested(self, route_id: int, entry_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        entry = next((e for e in route.entries if e.id == entry_id), None)
        if not entry:
            return
        if entry.visit_id and entry.visit_id in self._visits:
            self._remove_entry_to_pool(route, entry)
        elif entry.is_office_instance:
            self._recalc.remove_entry_from_route(route, entry, replace_with_empty=False)
            self._route_scene.rebuild_route(route)
            self._autosave.mark_dirty(route.id)

    # ------------------------------------------------------------------
    # Entry reordering
    # ------------------------------------------------------------------

    @Slot(int, int, int)
    def _on_entry_moved(self, route_id: int, entry_id: int, direction: int):
        route = self._routes.get(route_id)
        if not route:
            return
        entries = route.sorted_entries()
        idx = next((i for i, e in enumerate(entries) if e.id == entry_id), None)
        if idx is None:
            return
        target = idx + direction
        if target < 0 or target >= len(entries):
            return
        self._recalc.swap_entries(route, entries[idx], entries[target])
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route_id)

    # ------------------------------------------------------------------
    # Time editing
    # ------------------------------------------------------------------

    @Slot(int, int)
    def _on_entry_time_edit(self, route_id: int, entry_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        entry = next((e for e in route.entries if e.id == entry_id), None)
        if not entry:
            return

        entries = route.sorted_entries()
        entry_idx = next((i for i, e in enumerate(entries) if e.id == entry_id), 0)

        # Calculate minimum allowed start (must keep one-minute integrity)
        min_start = 0
        if entry_idx > 0:
            prev = entries[entry_idx - 1]
            seg = route.travel_segment_between(prev.id, entry.id)
            travel = seg.travel_minutes if seg else 0
            extra = self._settings.extra_time_minutes if route.extra_time_for_entry(entry.id) else 0
            min_start = _t2m(prev.end_time) + travel + extra

        current_start_abs = _t2m(entry.start_time)
        current_duration = max(1, _t2m(entry.end_time) - _t2m(entry.start_time))

        dlg = QDialog(self)
        dlg.setWindowTitle("Redigera besök")
        form = QFormLayout(dlg)

        start_edit = QLineEdit(_display_time(entry.start_time), dlg)
        start_edit.setPlaceholderText("HH:MM")
        duration_spin = QSpinBox(dlg)
        duration_spin.setRange(1, 720)
        duration_spin.setValue(current_duration)
        duration_spin.setSuffix(" min")
        office_name_edit: Optional[QLineEdit] = None
        office_address_edit: Optional[QLineEdit] = None

        if entry.is_office_instance:
            office_name_edit = QLineEdit(entry.office_name or "Kontor", dlg)
            office_address_edit = QLineEdit(entry.office_address or "", dlg)
            form.addRow("Namn:", office_name_edit)
            form.addRow("Adress:", office_address_edit)

        form.addRow("Starttid (HH:MM):", start_edit)
        form.addRow("Duration:", duration_spin)
        if entry_idx > 0:
            form.addRow("Tidigast:", QLabel(_display_time(_m2t(min_start)), dlg))

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel, parent=dlg)
        form.addRow(buttons)
        buttons.accepted.connect(dlg.accept)
        buttons.rejected.connect(dlg.reject)

        if dlg.exec() != QDialog.DialogCode.Accepted:
            return

        office_identity_changed = False
        pending_office_name = ""
        pending_office_address = ""
        if entry.is_office_instance:
            pending_office_name = (office_name_edit.text() if office_name_edit else "").strip() or "Kontor"
            pending_office_address = (office_address_edit.text() if office_address_edit else "").strip()
            if not pending_office_address:
                QMessageBox.warning(self, "Ogiltig adress", "Adress kan inte vara tom för standardbesök.")
                return
            office_identity_changed = (
                pending_office_name != (entry.office_name or "") or
                pending_office_address != (entry.office_address or "")
            )

        import re
        text = start_edit.text().strip()
        m = re.match(r"^(\d{1,2}):(\d{2})$", text)
        if not m:
            QMessageBox.warning(self, "Fel format", "Ange tid som HH:MM")
            return

        h = int(m.group(1))
        mm = int(m.group(2))
        if h < 0 or h > 23 or mm < 0 or mm > 59:
            QMessageBox.warning(self, "Fel format", "Ange giltig tid (00:00–23:59)")
            return

        target_day = current_start_abs // 1440
        new_start = target_day * 1440 + h * 60 + mm
        if current_start_abs - new_start > 720:
            new_start += 1440
        elif new_start - current_start_abs > 720:
            new_start -= 1440

        if new_start < min_start:
            QMessageBox.warning(
                self, "Ogiltig tid",
                f"Starttiden kräver tillräcklig lucka före besöket. Tidigast {_display_time(_m2t(min_start))}."
            )
            return

        duration = duration_spin.value()
        old_start = _t2m(entry.start_time)
        old_end = _t2m(entry.end_time)
        if entry.is_office_instance:
            entry.office_name = pending_office_name
            entry.office_address = pending_office_address
        entry.start_time = _m2t(new_start)
        entry.end_time = _m2t(new_start + duration)

        if entry_idx > 0:
            prev = entries[entry_idx - 1]
            seg = route.travel_segment_between(prev.id, entry.id)
            travel = seg.travel_minutes if seg else 0
            extra = self._settings.extra_time_minutes if route.extra_time_for_entry(entry.id) else 0
            base_start = _t2m(prev.end_time) + travel + extra
            manual_gap = max(0, new_start - base_start)
            space = route.empty_space_between(prev.id, entry.id)
            if manual_gap > 0:
                if space is None:
                    route.empty_spaces.append(EmptySpace(
                        id=None,
                        route_id=route.id,
                        from_entry_id=prev.id,
                        to_entry_id=entry.id,
                        duration_minutes=manual_gap,
                    ))
                else:
                    space.duration_minutes = manual_gap
            elif space is not None:
                route.empty_spaces.remove(space)

        new_end = _t2m(entry.end_time)
        self._recalc.shift_following_entries(route, entry.id, new_end - old_end,
                             include_anchor=False)

        self._recalc.recalculate(route)
        pair_changed_route_ids: set[int] = set()
        if entry.visit_id and self._pairs.get(entry.visit_id):
            pair_changed_route_ids = self._align_pair_to_start(entry.visit_id, new_start)

        rebuild_route_ids = {route.id}
        rebuild_route_ids.update(pair_changed_route_ids)
        for rebuild_route_id in rebuild_route_ids:
            rebuild_route = self._routes.get(rebuild_route_id)
            if rebuild_route:
                self._route_scene.rebuild_route(rebuild_route)
        if entry.is_office_instance:
            self._db.update_route_entry(entry)
            if office_identity_changed:
                self._request_travel_for_new_entry(route, entry)
        for dirty_route_id in rebuild_route_ids:
            self._autosave.mark_dirty(dirty_route_id)
        self._apply_pair_synchronization_and_refresh(max_passes=6)

    @Slot(int, int, int)
    def _on_entry_duration_changed(self, route_id: int, entry_id: int, delta: int):
        route = self._routes.get(route_id)
        if not route:
            return
        entry = next((e for e in route.entries if e.id == entry_id), None)
        if not entry:
            return
        current = _t2m(entry.end_time) - _t2m(entry.start_time)
        new_dur = max(1, current + delta)
        effective_delta = new_dur - current

        shift_delta = effective_delta
        if effective_delta > 0:
            entries = route.sorted_entries()
            idx = next((i for i, e in enumerate(entries) if e.id == entry.id), None)
            if idx is not None and idx < len(entries) - 1:
                nxt = entries[idx + 1]
                space = route.empty_space_between(entry.id, nxt.id)
                if space and space.duration_minutes > 0:
                    consume = min(effective_delta, space.duration_minutes)
                    space.duration_minutes -= consume
                    if space.duration_minutes <= 0:
                        route.empty_spaces.remove(space)
                    shift_delta = effective_delta - consume

        entry.end_time = _m2t(_t2m(entry.start_time) + new_dur)
        self._recalc.shift_following_entries(route, entry.id, shift_delta)
        self._recalc.recalculate(route)
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route.id)
        self._apply_pair_synchronization_and_refresh(max_passes=6)

    @Slot(int, int, int)
    def _on_empty_remove(self, route_id: int, from_entry_id: int, to_entry_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        space = route.empty_space_between(from_entry_id, to_entry_id)
        if not space:
            return
        removed = max(0, space.duration_minutes)
        route.empty_spaces.remove(space)
        to_entry = next((e for e in route.entries if e.id == to_entry_id), None)
        if to_entry and removed > 0:
            self._recalc.shift_following_entries(route, to_entry.id, -removed,
                                                 include_anchor=True)
        self._recalc.recalculate(route)
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route.id)

    # ------------------------------------------------------------------
    # Travel segment editing
    # ------------------------------------------------------------------

    @Slot(int, int, str)
    def _on_travel_mode_changed(self, route_id: int, seg_id: int, mode: str):
        route = self._routes.get(route_id)
        if not route:
            return
        seg = next((s for s in route.travel_segments if s.id == seg_id), None)
        if not seg:
            return
        previous_state = seg.travel_time_state
        previous_is_custom = seg.is_custom
        was_calculated_display = self._segment_is_calculated_display(seg)
        seg.mode = mode
        seg.api_failed = False
        seg.api_error = ""
        seg.is_calculating = False

        if previous_state == TravelTimeState.DEFAULT and not previous_is_custom:
            self._apply_segment_minutes(
                route,
                seg,
                self._settings.default_travel_for_mode(mode),
                is_custom=False,
                state=TravelTimeState.DEFAULT,
            )
            self._apply_pair_synchronization_and_refresh(max_passes=6)
            return

        if previous_state == TravelTimeState.CALCULATED and not previous_is_custom:
            seg.travel_time_state = TravelTimeState.CALCULATED
            self._recalc.recalculate(route)
            self._route_scene.rebuild_route(route)
            self._autosave.mark_dirty(route_id)
            self._apply_pair_synchronization_and_refresh(max_passes=6)
            self._request_segment_travel(route, seg, debounce_ms=700,
                                         display_is_calc=was_calculated_display)
            return

        seg.travel_time_state = TravelTimeState.EDITED if previous_is_custom else previous_state
        self._recalc.recalculate(route)
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route_id)
        self._apply_pair_synchronization_and_refresh(max_passes=6)

    @Slot(int, int, int)
    def _on_travel_duration_changed(self, route_id: int, seg_id: int, delta: int):
        route = self._routes.get(route_id)
        if not route:
            return
        seg = next((s for s in route.travel_segments if s.id == seg_id), None)
        if not seg:
            return
        old_minutes = max(0, seg.travel_minutes)
        new_minutes = max(0, old_minutes + delta)
        change = new_minutes - old_minutes

        if not seg.is_custom:
            seg.calculated_minutes = old_minutes
        seg.travel_minutes = new_minutes
        seg.is_custom = True
        seg.travel_time_state = TravelTimeState.EDITED
        seg.api_failed = False
        seg.api_error = ""
        seg.is_calculating = False

        shift_delta = change
        if change > 0:
            space = route.empty_space_between(seg.from_entry_id, seg.to_entry_id)
            if space and space.duration_minutes > 0:
                consume = min(change, space.duration_minutes)
                space.duration_minutes -= consume
                if space.duration_minutes <= 0:
                    route.empty_spaces.remove(space)
                shift_delta = change - consume

        if shift_delta != 0:
            to_entry = next((e for e in route.entries if e.id == seg.to_entry_id), None)
            if to_entry:
                self._recalc.shift_following_entries(route, to_entry.id, shift_delta,
                                                     include_anchor=True)

        self._recalc.recalculate(route)
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route_id)
        self._apply_pair_synchronization_and_refresh(max_passes=6)

    @Slot(int, int)
    def _on_travel_minutes_edit(self, route_id: int, seg_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        seg = next((s for s in route.travel_segments if s.id == seg_id), None)
        if not seg:
            return
        val, ok = QInputDialog.getInt(
            self, "Redigera resetid",
            "Resetid (minuter):",
            value=seg.travel_minutes, minValue=0, maxValue=300,
        )
        if ok:
            old_minutes = max(0, seg.travel_minutes)
            new_minutes = max(0, val)
            change = new_minutes - old_minutes

            if not seg.is_custom:
                seg.calculated_minutes = old_minutes
            seg.travel_minutes = new_minutes
            seg.is_custom = True
            seg.travel_time_state = TravelTimeState.EDITED
            seg.api_failed = False
            seg.api_error = ""
            seg.is_calculating = False

            shift_delta = change
            if change > 0:
                space = route.empty_space_between(seg.from_entry_id, seg.to_entry_id)
                if space and space.duration_minutes > 0:
                    consume = min(change, space.duration_minutes)
                    space.duration_minutes -= consume
                    if space.duration_minutes <= 0:
                        route.empty_spaces.remove(space)
                    shift_delta = change - consume

            if shift_delta != 0:
                to_entry = next((e for e in route.entries if e.id == seg.to_entry_id), None)
                if to_entry:
                    self._recalc.shift_following_entries(route, to_entry.id, shift_delta,
                                                         include_anchor=True)

            self._recalc.recalculate(route)
            self._route_scene.rebuild_route(route)
            self._autosave.mark_dirty(route_id)
            self._apply_pair_synchronization_and_refresh(max_passes=6)

    @Slot(int, int)
    def _on_travel_retry(self, route_id: int, seg_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        seg = next((s for s in route.travel_segments if s.id == seg_id), None)
        if not seg:
            return
        self._request_segment_travel(route, seg, debounce_ms=0,
                                     display_is_calc=self._segment_is_calculated_display(seg))

    @Slot(int, int)
    def _on_travel_source_toggle(self, route_id: int, seg_id: int):
        route = self._routes.get(route_id)
        if not route:
            return
        seg = next((s for s in route.travel_segments if s.id == seg_id), None)
        if not seg:
            return

        if seg.travel_time_state == TravelTimeState.EDITED:
            self._apply_segment_minutes(route, seg,
                                        self._settings.default_travel_for_mode(seg.mode),
                                        is_custom=False,
                                        state=TravelTimeState.DEFAULT)
            return

        if seg.travel_time_state == TravelTimeState.CALCULATED and not seg.api_failed:
            self._apply_segment_minutes(route, seg,
                                        self._settings.default_travel_for_mode(seg.mode),
                                        is_custom=False,
                                        state=TravelTimeState.DEFAULT)
            return

        self._request_segment_travel(route, seg, debounce_ms=0,
                         display_is_calc=self._segment_is_calculated_display(seg))

    def _apply_segment_minutes(self, route: Route, seg: TravelSegment, minutes: int,
                               is_custom: bool, state: str):
        old_minutes = max(0, seg.travel_minutes)
        new_minutes = max(0, int(minutes))
        change = new_minutes - old_minutes

        seg.travel_minutes = new_minutes
        seg.is_custom = is_custom
        seg.travel_time_state = state
        seg.api_failed = False
        seg.api_error = ""
        seg.is_calculating = False

        shift_delta = change
        if change > 0:
            space = route.empty_space_between(seg.from_entry_id, seg.to_entry_id)
            if space and space.duration_minutes > 0:
                consume = min(change, space.duration_minutes)
                space.duration_minutes -= consume
                if space.duration_minutes <= 0:
                    route.empty_spaces.remove(space)
                shift_delta = change - consume

        if shift_delta != 0:
            to_entry = next((e for e in route.entries if e.id == seg.to_entry_id), None)
            if to_entry:
                self._recalc.shift_following_entries(route, to_entry.id, shift_delta,
                                                     include_anchor=True)

        self._recalc.recalculate(route)
        self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty(route.id)

    # ------------------------------------------------------------------
    # Travel time async result
    # ------------------------------------------------------------------

    @Slot(str, str, str, int)
    def _on_travel_time_ready(self, from_addr: str, to_addr: str,
                               mode: str, minutes: int):
        """Update any route segments that use this pair."""
        key = (from_addr, to_addr, mode)
        from_failed_fallback = key in self._failed_fallback_keys
        if from_failed_fallback:
            self._failed_fallback_keys.discard(key)

        any_route_changed = False

        for route in self._routes.values():
            changed = False
            for seg in route.travel_segments:
                if seg.is_custom:
                    continue
                if not seg.is_calculating:
                    continue
                from_entry = next(
                    (e for e in route.entries if e.id == seg.from_entry_id), None)
                to_entry = next(
                    (e for e in route.entries if e.id == seg.to_entry_id), None)
                if not from_entry or not to_entry:
                    continue
                if (from_entry.api_address == from_addr and
                        to_entry.api_address == to_addr and
                        seg.mode == mode):
                    seg.travel_minutes = minutes
                    seg.calculated_minutes = minutes
                    seg.is_calculating = False
                    seg.loading_display_is_calc = None
                    if not from_failed_fallback:
                        seg.api_failed = False
                        seg.api_error = ""
                        seg.travel_time_state = TravelTimeState.CALCULATED
                    changed = True
            if changed:
                any_route_changed = True
                self._recalc.recalculate(route)
                self._route_scene.rebuild_route(route)
                self._autosave.mark_dirty(route.id)
        if any_route_changed:
            self._apply_pair_synchronization_and_refresh(max_passes=6)
        if self._settings.debug_mode:
            self._run_debug_integrity_scan("travel_time_ready", log_result=False)

    def _connect_travel_signals(self):
        self._travel_svc.travel_time_ready.connect(self._on_travel_time_ready)
        self._travel_svc.lookup_started.connect(self._on_travel_lookup_started)
        self._travel_svc.request_logged.connect(
            lambda t: self._get_travel_status().log_request(t))
        self._travel_svc.response_logged.connect(
            lambda t: self._get_travel_status().log_response(t))
        self._travel_svc.error_occurred.connect(self._on_travel_time_error)
        self._travel_svc.quota_warning.connect(
            lambda c: (self._get_travel_status().log_quota_warning(c),
                       QMessageBox.warning(
                           self, "API-kvot",
                           f"Du har nu gjort {c} API-anrop. "
                           "Kontrollera ditt kvotgräns i Inställningar."
                       )))

    def _request_travel_for_new_entry(self, route: Route, new_entry: RouteEntry):
        """Fire async travel time requests for neighbors of a newly placed entry."""
        entries = route.sorted_entries()
        idx = next((i for i, e in enumerate(entries) if e.id == new_entry.id), None)
        if idx is None:
            return
        if idx > 0:
            prev = entries[idx - 1]
            seg_prev = route.travel_segment_between(prev.id, new_entry.id)
            if seg_prev:
                self._request_segment_travel(route, seg_prev, debounce_ms=0)
        if idx < len(entries) - 1:
            nxt = entries[idx + 1]
            seg_next = route.travel_segment_between(new_entry.id, nxt.id)
            if seg_next:
                self._request_segment_travel(route, seg_next, debounce_ms=0)

    @Slot(str, str, str)
    def _on_travel_lookup_started(self, from_addr: str, to_addr: str, mode: str):
        self._set_segments_lookup_state(from_addr, to_addr, mode, is_calculating=True,
                                        api_failed=False, api_error="")
        if self._settings.debug_mode:
            self._run_debug_integrity_scan("travel_lookup_started", log_result=False)

    @Slot(str, str, str, str)
    def _on_travel_time_error(self, from_addr: str, to_addr: str, mode: str, error_msg: str):
        self._get_travel_status().log_error(from_addr, to_addr, mode, error_msg)
        self._failed_fallback_keys.add((from_addr, to_addr, mode))
        self._set_segments_lookup_state(from_addr, to_addr, mode, is_calculating=False,
                                        api_failed=True, api_error=error_msg)
        if self._settings.debug_mode:
            self._run_debug_integrity_scan("travel_time_error", log_result=False)

    def _set_segments_lookup_state(self, from_addr: str, to_addr: str, mode: str,
                                   is_calculating: bool, api_failed: bool,
                                   api_error: str = ""):
        for route in self._routes.values():
            changed = False
            for seg in route.travel_segments:
                from_entry = next((e for e in route.entries if e.id == seg.from_entry_id), None)
                to_entry = next((e for e in route.entries if e.id == seg.to_entry_id), None)
                if not from_entry or not to_entry:
                    continue
                if (from_entry.api_address == from_addr and
                        to_entry.api_address == to_addr and
                        seg.mode == mode):
                    if is_calculating and seg.loading_display_is_calc is None:
                        seg.loading_display_is_calc = self._segment_is_calculated_display(seg)
                    seg.is_calculating = is_calculating
                    seg.api_failed = api_failed
                    seg.api_error = api_error
                    if not is_calculating:
                        seg.loading_display_is_calc = None
                    if api_failed:
                        seg.travel_time_state = TravelTimeState.DEFAULT
                    changed = True
            if changed:
                self._route_scene.rebuild_route(route)

    def _request_segment_travel(self, route: Route, seg: TravelSegment,
                                debounce_ms: int = 0,
                                display_is_calc: Optional[bool] = None):
        from_entry = next((e for e in route.entries if e.id == seg.from_entry_id), None)
        to_entry = next((e for e in route.entries if e.id == seg.to_entry_id), None)
        if not from_entry or not to_entry:
            return

        from_addr = from_entry.api_address
        to_addr = to_entry.api_address
        if not from_addr or not to_addr:
            return

        if display_is_calc is None:
            display_is_calc = self._segment_is_calculated_display(seg)
        seg.loading_display_is_calc = bool(display_is_calc)

        seg.is_calculating = True
        seg.api_failed = False
        seg.api_error = ""
        self._route_scene.rebuild_route(route)

        timer_key = (route.id, seg.from_entry_id, seg.to_entry_id)
        existing = self._travel_debounce_timers.pop(timer_key, None)
        if existing is not None:
            existing.stop()
            existing.deleteLater()

        def _fire_request():
            self._travel_debounce_timers.pop(timer_key, None)
            self._travel_svc.request_travel_time(from_addr, to_addr, seg.mode)

        if debounce_ms > 0:
            timer = QTimer(self)
            timer.setSingleShot(True)
            timer.timeout.connect(_fire_request)
            self._travel_debounce_timers[timer_key] = timer
            timer.start(int(debounce_ms))
        else:
            _fire_request()

    def _segment_is_calculated_display(self, seg: TravelSegment) -> bool:
        return (
            seg.travel_time_state == TravelTimeState.CALCULATED
            and not seg.is_custom
            and not seg.api_failed
        )

    @Slot()
    def _on_time_integrity(self):
        before = self._capture_integrity_snapshot()
        for route in self._routes.values():
            self._recalc.recalculate(route)
            self._route_scene.rebuild_route(route)
        after = self._capture_integrity_snapshot()
        self._log_integrity_changes(before, after)
        self._autosave.flush_now()

    def _capture_integrity_snapshot(self) -> dict:
        snapshot: dict = {
            "entries": {},
            "segments": {},
            "spaces": {},
        }
        for route in self._routes.values():
            for entry in route.entries:
                if entry.id is None:
                    continue
                snapshot["entries"][entry.id] = (
                    route.id,
                    route.name,
                    entry.display_name,
                    entry.start_time,
                    entry.end_time,
                )
            for seg in route.travel_segments:
                if seg.id is None:
                    continue
                snapshot["segments"][seg.id] = (
                    route.id,
                    route.name,
                    seg.from_entry_id,
                    seg.to_entry_id,
                    seg.mode,
                    seg.travel_minutes,
                    seg.start_time,
                    seg.end_time,
                )
            for esp in route.empty_spaces:
                if esp.id is None:
                    continue
                snapshot["spaces"][esp.id] = (
                    route.id,
                    route.name,
                    esp.from_entry_id,
                    esp.to_entry_id,
                    esp.duration_minutes,
                    esp.start_time,
                    esp.end_time,
                )
        return snapshot

    def _log_integrity_changes(self, before: dict, after: dict):
        lines: list[str] = []

        for entry_id, after_val in after["entries"].items():
            before_val = before["entries"].get(entry_id)
            if before_val and before_val[3:] != after_val[3:]:
                route_name = after_val[1]
                visit_name = after_val[2]
                lines.append(
                    f"{route_name} | Visit '{visit_name}': "
                    f"{_display_time(before_val[3])}–{_display_time(before_val[4])} → "
                    f"{_display_time(after_val[3])}–{_display_time(after_val[4])}"
                )

        for seg_id, after_val in after["segments"].items():
            before_val = before["segments"].get(seg_id)
            if before_val and before_val[5:] != after_val[5:]:
                route_name = after_val[1]
                lines.append(
                    f"{route_name} | Travel {after_val[2]}→{after_val[3]}: "
                    f"{before_val[5]} min ({_display_time(before_val[6])}–{_display_time(before_val[7])}) → "
                    f"{after_val[5]} min ({_display_time(after_val[6])}–{_display_time(after_val[7])})"
                )

        for space_id, after_val in after["spaces"].items():
            before_val = before["spaces"].get(space_id)
            if before_val and before_val[4:] != after_val[4:]:
                route_name = after_val[1]
                lines.append(
                    f"{route_name} | Gap {after_val[2]}→{after_val[3]}: "
                    f"{before_val[4]} min ({_display_time(before_val[5])}–{_display_time(before_val[6])}) → "
                    f"{after_val[4]} min ({_display_time(after_val[5])}–{_display_time(after_val[6])})"
                )

        if lines:
            self._get_travel_status().log_integrity_change(
                f"Time integrity made {len(lines)} changes"
            )
            for line in lines[:120]:
                self._get_travel_status().log_integrity_change(line)
            if len(lines) > 120:
                self._get_travel_status().log_integrity_change(
                    f"... and {len(lines) - 120} more"
                )
        else:
            self._get_travel_status().log_integrity_ok("Time integrity: no changes needed")

    @Slot()
    def _on_block_visibility_changed(self):
        self._settings.show_travel_blocks = self._show_travel_action.isChecked()
        self._settings.show_space_blocks = self._show_space_action.isChecked()
        self._settings.show_extra_time_blocks = self._show_extra_time_action.isChecked()
        self._route_scene.set_block_visibility(
            self._settings.show_travel_blocks,
            self._settings.show_space_blocks,
            self._settings.show_extra_time_blocks,
        )
        self._autosave.mark_dirty()

    @Slot(bool)
    def _on_pair_alignment_toggled(self, checked: bool):
        self._settings.align_pair_visits = bool(checked)
        self._route_scene.set_pair_alignment_enabled(self._settings.align_pair_visits)
        if not self._settings.align_pair_visits:
            self._last_pair_alignment_conflict_count = 0
            self.statusBar().clearMessage()
        self._autosave.mark_dirty()

    @Slot(int)
    def _on_pair_alignment_conflicts_changed(self, count: int):
        conflict_count = max(0, int(count))
        if conflict_count == self._last_pair_alignment_conflict_count:
            return
        self._last_pair_alignment_conflict_count = conflict_count
        if not self._settings.align_pair_visits:
            return
        if conflict_count <= 0:
            self.statusBar().showMessage("Par-justering: alla par kunde justeras.", 2500)
            return
        if conflict_count == 1:
            self.statusBar().showMessage(
                "Par-justering: 1 par hoppades över (korsande ordning mellan rutter).",
                7000,
            )
            return
        self.statusBar().showMessage(
            f"Par-justering: {conflict_count} par hoppades över (korsande ordning mellan rutter).",
            7000,
        )

    @Slot(bool)
    def _on_extra_time_auto_toggled(self, checked: bool):
        self._settings.extra_time_auto_place = bool(checked)
        self._autosave.mark_dirty()

    @Slot(int)
    def _on_extra_time_minutes_changed(self, minutes: int):
        self._settings.extra_time_minutes = max(0, int(minutes))
        if self._default_templates and self._default_templates[-1].get("type") == "extra_time":
            self._default_templates[-1]["duration_minutes"] = self._settings.extra_time_minutes
        self._pool_scene.set_extra_time_minutes(self._settings.extra_time_minutes)
        self._route_scene.set_extra_time_minutes(self._settings.extra_time_minutes)
        for route in self._routes.values():
            self._recalc.recalculate(route)
            self._route_scene.rebuild_route(route)
        self._autosave.mark_dirty()

    @Slot()
    def _on_reset_all(self):
        reply = QMessageBox.question(
            self, "Rensa allt",
            "Detta tar bort alla rutter och besök från planeringen. Fortsätta?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        current_state = self._db.export_state()
        self._db.import_state({
            "visits": [],
            "routes": [],
            "route_visit_order": [],
            "travel_segments": [],
            "empty_spaces": [],
            "extra_time_blocks": [],
            "settings": current_state.get("settings", []),
            "column_order": current_state.get("column_order", []),
            "office_template": current_state.get("office_template", []),
        })
        self._load_data()
        self._autosave.clear()
        QMessageBox.information(self, "Rensat", "All planering har rensats.")

    # ------------------------------------------------------------------
    # Default travel mode
    # ------------------------------------------------------------------

    @Slot(int)
    def _on_default_mode_changed(self, idx: int):
        self._settings.default_travel_mode = self._mode_combo.currentData()
        self._autosave.mark_dirty()

    # ------------------------------------------------------------------
    # UI scale
    # ------------------------------------------------------------------

    @Slot(int)
    def _on_ui_scale_percent_changed(self, percent: int):
        clamped_percent = max(50, min(200, int(percent)))
        self._settings.ui_scale_percent = clamped_percent
        scaled_font_size = max(6, min(30, int(round(12 * clamped_percent / 100.0))))
        self._settings.font_size = scaled_font_size
        self._layout_engine.font_size = scaled_font_size
        self._route_scene.set_font_size(scaled_font_size)
        self._pool_scene.set_font_size(scaled_font_size)
        self._autosave.mark_dirty()

    # ------------------------------------------------------------------
    # Excel import
    # ------------------------------------------------------------------

    @Slot()
    def _on_import_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Öppna Excel-fil", "", "Excel-filer (*.xls *.xlsx)"
        )
        if not path:
            return

        try:
            visits, result = self._import_svc.preview_import(path)
        except Exception as exc:
            QMessageBox.critical(self, "Importfel", str(exc))
            return

        if result.duplicates_in_file:
            QMessageBox.warning(
                self, "Dubbletter i filen",
                "Följande ObjectID förekommer flera gånger i filen:\n"
                + "\n".join(result.duplicates_in_file[:20]),
            )

        msg = (f"Filen innehåller {result.total_rows} rader.\n"
               f"Att importera: {result.imported}  |  Hoppar över: {result.skipped}")
        reply = QMessageBox.question(
            self, "Bekräfta import", msg,
            QMessageBox.StandardButton.Ok | QMessageBox.StandardButton.Cancel,
        )
        if reply != QMessageBox.StandardButton.Ok:
            return

        removed_ids = self._import_svc.commit_import(visits, path, result)

        if removed_ids:
            reply2 = QMessageBox.question(
                self, "Besök som saknas i importen",
                f"{len(removed_ids)} besök finns i databasen men inte i den nya filen.\n"
                "Vill du ta bort dem? (De ersätts med tomrum i rutterna.)",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply2 == QMessageBox.StandardButton.Yes:
                self._import_svc.remove_visits_not_in_import(removed_ids)
                # Rebuild routes that had those visits
                routes = self._db.load_all_routes(
                    {v.id: v for v in self._db.load_all_visits()})
                for r in routes:
                    self._routes[r.id] = r
                    self._route_scene.rebuild_route(r)

        # Reload everything
        self._load_data()
        QMessageBox.information(
            self, "Import klar",
            f"Importerade {result.imported} besök."
        )

    # ------------------------------------------------------------------
    # Settings dialog
    # ------------------------------------------------------------------

    @Slot()
    def _on_open_settings(self):
        dlg = SettingsDialog(self._settings, self._api_key, self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            previous_debug_mode = bool(self._settings.debug_mode)
            self._settings = dlg.get_settings()
            self._apply_visit_color_palette()
            self._recalc.set_settings(self._settings)
            new_key = dlg.get_api_key()
            if new_key != self._api_key:
                self._api_key = new_key
                _save_api_key(new_key)
                self._travel_svc.set_api_key(new_key)
                if self._map_window is not None:
                    self._map_window.close()
                    self._map_window = None
            self._db.save_settings(self._settings)
            # Sync toolbar widgets
            self._ui_scale_spin.blockSignals(True)
            self._ui_scale_spin.setValue(max(50, min(200, int(self._settings.ui_scale_percent))))
            self._ui_scale_spin.blockSignals(False)
            self._on_ui_scale_percent_changed(self._settings.ui_scale_percent)
            idx = self._mode_combo.findData(self._settings.default_travel_mode)
            if idx >= 0:
                self._mode_combo.setCurrentIndex(idx)
            self._extra_time_auto_action.setChecked(bool(self._settings.extra_time_auto_place))
            self._extra_time_spin.setValue(max(0, int(self._settings.extra_time_minutes)))
            self._pair_align_action.setChecked(bool(self._settings.align_pair_visits))
            self._get_travel_status().configure_file_logging(
                self._settings.file_logging_enabled,
                self._settings.file_logging_retention_days,
            )

            if self._settings.debug_mode:
                if not previous_debug_mode:
                    self._get_travel_status().log_debug("Debug mode enabled")
                self._run_debug_integrity_scan("Settings")
            elif previous_debug_mode:
                self._get_travel_status().log_debug("Debug mode disabled")
                self._clear_debug_integrity_marks()

    def _apply_visit_color_palette(self):
        VisitItem.set_color_palette(self._settings.visit_ribbon_color_map())
        for scene in (getattr(self, "_route_scene", None), getattr(self, "_pool_scene", None)):
            if scene is None:
                continue
            for item in scene.items():
                if isinstance(item, VisitItem):
                    item.update()

    # ------------------------------------------------------------------
    # State export / import
    # ------------------------------------------------------------------

    @Slot()
    def _on_export_state(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Spara tillstånd", "planning_state.json",
            "JSON-filer (*.json)"
        )
        if not path:
            return
        state = self._db.export_state()
        with open(path, "w", encoding="utf-8") as f:
            json.dump(state, f, ensure_ascii=False, indent=2)
        QMessageBox.information(self, "Export klar", f"Tillståndet sparades till\n{path}")

    @Slot()
    def _on_import_state(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Ladda tillstånd", "", "JSON-filer (*.json)"
        )
        if not path:
            return
        reply = QMessageBox.question(
            self, "Bekräfta",
            "Detta ersätter all nuvarande data (besök och rutter). Fortsätta?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                state = json.load(f)
            self._db.import_state(state)
            self._load_data()
            QMessageBox.information(self, "Import klar", "Tillståndet laddades.")
        except Exception as exc:
            QMessageBox.critical(self, "Fel vid import", str(exc))

    # ------------------------------------------------------------------
    # Excel export
    # ------------------------------------------------------------------

    @Slot()
    def _on_export_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Exportera Excel", "export.xlsx",
            "Excel-filer (*.xlsx)"
        )
        if not path:
            return
        try:
            self._write_excel_export(path)
            QMessageBox.information(self, "Export klar", f"Exporterat till\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Exportfel", str(exc))

    @Slot()
    def _on_export_pdf(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Exportera PDF", "export.pdf",
            "PDF-filer (*.pdf)"
        )
        if not path:
            return
        try:
            self._write_pdf_export(path)
            QMessageBox.information(self, "Export klar", f"Exporterat till\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Exportfel", str(exc))

    def _to_excel_sheet_name(self, route_name: str, used_names: set[str]) -> str:
        base = re.sub(r"[\\/*?:\[\]]", "_", (route_name or "").strip())
        if not base:
            base = "Rutt"
        base = base[:31]
        candidate = base
        suffix_num = 2
        while candidate in used_names:
            suffix = f" ({suffix_num})"
            candidate = f"{base[:31 - len(suffix)]}{suffix}"
            suffix_num += 1
        used_names.add(candidate)
        return candidate

    def _write_excel_export(self, path: str):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()

        # Remove default sheet so only route sheets are exported.
        wb.remove(wb.active)

        header_labels = ["Start", "Slut", "Namn", "Adress", "Insatser", "Färdsätt"]
        summary_header_labels = ["Rutt", "Start", "Slut", "Namn", "Adress", "Insatser", "Färdsätt"]
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor="4F81BD")
        odd_fill = PatternFill(fill_type="solid", fgColor="DCE6F1")
        even_fill = PatternFill(fill_type="solid", fgColor="EDF2F9")

        def _mode_label(mode: Optional[str]) -> str:
            if mode == TravelMode.CAR:
                return "Bil"
            if mode == TravelMode.BIKE:
                return "Cykel"
            if mode == TravelMode.WALK:
                return "Gå"
            return "-"

        summary_rows: list[list[str]] = []

        used_sheet_names: set[str] = set()
        routes = sorted(self._routes.values(), key=lambda r: r.display_order)

        for route in routes:
            sheet_name = self._to_excel_sheet_name(route.name, used_sheet_names)
            ws = wb.create_sheet(sheet_name)
            ws.append(header_labels)

            for col_idx in range(1, len(header_labels) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill

            entries = route.sorted_entries()
            for data_idx, entry in enumerate(entries, start=1):
                mode_label = "-"
                if data_idx > 1:
                    prev = entries[data_idx - 2]
                    seg = route.travel_segment_between(prev.id, entry.id)
                    mode_label = _mode_label(seg.mode if seg else None)

                ws.append([
                    _display_time(entry.start_time),
                    _display_time(entry.end_time),
                    entry.display_name,
                    entry.display_address,
                    entry.display_insatser,
                    mode_label,
                ])
                summary_rows.append([
                    route.name,
                    _display_time(entry.start_time),
                    _display_time(entry.end_time),
                    entry.display_name,
                    entry.display_address,
                    entry.display_insatser,
                    mode_label,
                ])
                row_idx = data_idx + 1
                row_fill = odd_fill if data_idx % 2 == 1 else even_fill
                for col_idx in range(1, len(header_labels) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = row_fill

            ws.auto_filter.ref = f"A1:F{max(1, ws.max_row)}"

            for col_idx in range(1, len(header_labels) + 1):
                max_len = 0
                for row_idx in range(1, ws.max_row + 1):
                    value = ws.cell(row=row_idx, column=col_idx).value
                    text = "" if value is None else str(value)
                    if len(text) > max_len:
                        max_len = len(text)
                ws.column_dimensions[get_column_letter(col_idx)].width = max(14, max_len + 2)

        if not routes:
            ws = wb.create_sheet("Rutter")
            ws.append(header_labels)
            for col_idx in range(1, len(header_labels) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
            ws.auto_filter.ref = "A1:F1"

        ws_summary = wb.create_sheet("Sammanfattning")
        ws_summary.append(summary_header_labels)
        for col_idx in range(1, len(summary_header_labels) + 1):
            cell = ws_summary.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill

        for data_idx, row_data in enumerate(summary_rows, start=1):
            ws_summary.append(row_data)
            row_idx = data_idx + 1
            row_fill = odd_fill if data_idx % 2 == 1 else even_fill
            for col_idx in range(1, len(summary_header_labels) + 1):
                ws_summary.cell(row=row_idx, column=col_idx).fill = row_fill

        ws_summary.auto_filter.ref = f"A1:G{max(1, ws_summary.max_row)}"
        for col_idx in range(1, len(summary_header_labels) + 1):
            max_len = 0
            for row_idx in range(1, ws_summary.max_row + 1):
                value = ws_summary.cell(row=row_idx, column=col_idx).value
                text = "" if value is None else str(value)
                if len(text) > max_len:
                    max_len = len(text)
            ws_summary.column_dimensions[get_column_letter(col_idx)].width = max(14, max_len + 2)

        wb.save(path)

    def _write_pdf_export(self, path: str):
        writer = QPdfWriter(path)
        writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        writer.setPageOrientation(QPageLayout.Orientation.Portrait)
        writer.setResolution(144)

        painter = QPainter()
        if not painter.begin(writer):
            raise RuntimeError("Kunde inte skapa PDF-export.")

        try:
            page_rect = writer.pageLayout().paintRectPixels(writer.resolution())
            margin = 36
            left = page_rect.left() + margin
            top = page_rect.top() + margin
            right = page_rect.right() - margin
            bottom = page_rect.bottom() - margin
            page_w = right - left
            page_h = bottom - top

            route_name_font = QFont("Segoe UI", 11, QFont.Weight.Bold)
            body_font = QFont("Segoe UI", 9)
            small_font = QFont("Segoe UI", 8)

            color_map = {
                key: QColor(value)
                for key, value in self._settings.visit_ribbon_color_map().items()
            }

            column_gap = 18
            columns_per_page = 2
            column_w = int((page_w - (column_gap * (columns_per_page - 1))) / columns_per_page)

            header_h = 98
            visit_h = 72
            travel_h = 44
            empty_h = 30
            extra_h = 30
            block_gap = 8
            content_pad = 10
            left_strip_w = 6

            routes = sorted(self._routes.values(), key=lambda r: r.display_order)
            if not routes:
                painter.setFont(route_name_font)
                painter.setPen(QPen(QColor("#000000")))
                painter.drawText(
                    QRectF(left, top, page_w, 40),
                    Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                    "Inga rutter att exportera",
                )
                return

            def _pdf_mode_label(mode: Optional[str]) -> str:
                if mode == TravelMode.CAR:
                    return "Bil"
                if mode == TravelMode.BIKE:
                    return "Cykel"
                if mode == TravelMode.WALK:
                    return "Gå"
                return "-"

            def draw_route_column(
                route: Route,
                x: int,
                y: int,
                items_subset: list[tuple[str, object]],
                continued: bool = False,
            ):
                col_h = page_h

                painter.fillRect(QRectF(x, y, column_w, col_h), QColor(COLOR_ROUTE_COLUMN_BG))
                painter.setPen(QPen(QColor("#B0BEC5"), 1))
                painter.drawRect(QRectF(x, y, column_w, col_h))

                painter.fillRect(QRectF(x, y, column_w, header_h), QColor(COLOR_HEADER_BG))
                painter.setPen(QPen(QColor("#CFD8DC"), 1))
                painter.drawLine(x, y + header_h, x + column_w, y + header_h)

                painter.setPen(QPen(QColor("#212121")))
                painter.setFont(route_name_font)
                painter.drawText(
                    QRectF(x + content_pad, y + 6, column_w - (content_pad * 2), 28),
                    Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                    f"{route.name or 'Rutt'}{' (forts.)' if continued else ''}",
                )

                painter.setFont(small_font)
                painter.setPen(QPen(QColor("#546E7A")))
                painter.drawText(
                    QRectF(x + content_pad, y + 36, column_w - (content_pad * 2), 16),
                    Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                    "Anteckningar:",
                )
                painter.setPen(QPen(QColor("#37474F")))
                painter.drawText(
                    QRectF(x + content_pad, y + 52, column_w - (content_pad * 2), 40),
                    Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignTop | Qt.TextFlag.TextWordWrap,
                    route.notes.strip() if route.notes.strip() else "-",
                )

                row_y = y + header_h + block_gap
                for item_type, item in items_subset:
                    if item_type == "travel":
                        seg = item
                        is_verklig = bool(
                            not seg.is_custom
                            and seg.calculated_minutes is not None
                            and not seg.api_failed
                        )
                        travel_bg = QColor(COLOR_TRAVEL_BG) if is_verklig else QColor(COLOR_TRAVEL_DEFAULT_BG)
                        travel_border = QColor(COLOR_TRAVEL_BORDER) if is_verklig else QColor(COLOR_TRAVEL_DEFAULT_BORDER)
                        travel_text = QColor("#5D4037") if is_verklig else QColor("#424242")
                        travel_value = QColor("#1565C0")

                        painter.fillRect(
                            QRectF(x + content_pad, row_y, column_w - (content_pad * 2), travel_h),
                            travel_bg,
                        )
                        painter.setPen(QPen(travel_border, 1))
                        painter.drawRect(QRectF(x + content_pad, row_y, column_w - (content_pad * 2), travel_h))

                        mode_label = {
                            TravelMode.CAR: "Bil",
                            TravelMode.BIKE: "Cykel",
                            TravelMode.WALK: "Gå",
                        }.get(seg.mode, seg.mode)
                        if seg.is_custom:
                            source_label = "Redigerad restid"
                        elif is_verklig:
                            source_label = "Verklig restid"
                        else:
                            source_label = "Standard restid"

                        painter.setFont(small_font)
                        painter.setPen(QPen(travel_text))
                        painter.drawText(
                            QRectF(x + content_pad + 8, row_y + 5, column_w - (content_pad * 2) - 90, 16),
                            Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                            source_label,
                        )
                        painter.drawText(
                            QRectF(x + content_pad + 8, row_y + 22, column_w - (content_pad * 2) - 90, 16),
                            Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                            f"Restid ({mode_label})",
                        )
                        painter.setFont(body_font)
                        painter.setPen(QPen(travel_value))
                        painter.drawText(
                            QRectF(x + column_w - content_pad - 80, row_y + 9, 72, 24),
                            Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                            f"{seg.travel_minutes} min",
                        )
                        row_y += travel_h + block_gap
                        continue

                    if item_type == "empty":
                        esp = item
                        painter.fillRect(
                            QRectF(x + content_pad, row_y, column_w - (content_pad * 2), empty_h),
                            QColor(COLOR_EMPTY_BG),
                        )
                        painter.setPen(QPen(QColor(COLOR_EMPTY_BORDER), 1))
                        painter.drawRect(QRectF(x + content_pad, row_y, column_w - (content_pad * 2), empty_h))
                        painter.setFont(small_font)
                        painter.setPen(QPen(QColor("#0D47A1")))
                        painter.drawText(
                            QRectF(x + content_pad + 8, row_y, column_w - (content_pad * 2) - 16, empty_h),
                            Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                            f"Lucka: {esp.duration_minutes} min",
                        )
                        row_y += empty_h + block_gap
                        continue

                    if item_type == "extra":
                        painter.fillRect(
                            QRectF(x + content_pad, row_y, column_w - (content_pad * 2), extra_h),
                            QColor("#BBDEFB"),
                        )
                        painter.setPen(QPen(QColor("#64B5F6"), 1))
                        painter.drawRect(QRectF(x + content_pad, row_y, column_w - (content_pad * 2), extra_h))
                        painter.setFont(small_font)
                        painter.setPen(QPen(QColor("#0D47A1")))
                        painter.drawText(
                            QRectF(x + content_pad + 8, row_y, column_w - (content_pad * 2) - 16, extra_h),
                            Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                            f"Extratid: {self._settings.extra_time_minutes} min",
                        )
                        row_y += extra_h + block_gap
                        continue

                    entry, incoming_mode_label = item
                    painter.fillRect(QRectF(x + content_pad, row_y, column_w - (content_pad * 2), visit_h), QColor(COLOR_VISIT_BG))

                    strip_color = color_map.get(entry.display_color or "", QColor("#2B2B2B"))
                    painter.fillRect(
                        QRectF(x + content_pad, row_y, left_strip_w, visit_h),
                        strip_color,
                    )

                    painter.setPen(QPen(QColor(COLOR_VISIT_BORDER), 1))
                    painter.drawRect(QRectF(x + content_pad, row_y, column_w - (content_pad * 2), visit_h))

                    text_x = x + content_pad + left_strip_w + 8
                    text_w = column_w - (content_pad * 2) - left_strip_w - 86

                    painter.setFont(body_font)
                    painter.setPen(QPen(QColor("#212121")))
                    painter.drawText(
                        QRectF(text_x, row_y + 6, text_w, 22),
                        Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                        entry.display_name,
                    )

                    painter.setFont(small_font)
                    painter.setPen(QPen(QColor("#424242")))
                    painter.drawText(
                        QRectF(text_x, row_y + 28, text_w, 18),
                        Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                        entry.display_address,
                    )
                    painter.drawText(
                        QRectF(text_x, row_y + 46, text_w, 16),
                        Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                        entry.display_insatser,
                    )
                    painter.drawText(
                        QRectF(text_x, row_y + 60, text_w, 12),
                        Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                        f"Färdsätt: {incoming_mode_label}",
                    )

                    painter.setFont(body_font)
                    painter.setPen(QPen(QColor("#0D47A1")))
                    visit_duration = max(0, _t2m(entry.end_time) - _t2m(entry.start_time))
                    painter.drawText(
                        QRectF(x + column_w - content_pad - 72, row_y + 8, 68, 22),
                        Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                        _display_time(entry.start_time),
                    )
                    painter.drawText(
                        QRectF(x + column_w - content_pad - 72, row_y + 34, 68, 22),
                        Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                        _display_time(entry.end_time),
                    )
                    painter.setFont(small_font)
                    painter.setPen(QPen(QColor("#1565C0")))
                    painter.drawText(
                        QRectF(x + column_w - content_pad - 72, row_y + 52, 68, 16),
                        Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                        f"{visit_duration} min",
                    )

                    row_y += visit_h + block_gap

            available_h = page_h - header_h - content_pad - block_gap

            route_chunk_map: dict[int, list[list[tuple[str, object]]]] = {}
            for route in routes:
                entries = route.sorted_entries()
                route_items: list[tuple[str, object]] = []
                for idx, entry in enumerate(entries):
                    incoming_mode_label = "-"
                    if idx > 0:
                        prev = entries[idx - 1]
                        if self._settings.show_extra_time_blocks and route.extra_time_for_entry(entry.id):
                            route_items.append(("extra", route.extra_time_for_entry(entry.id)))
                        if self._settings.show_travel_blocks:
                            seg = route.travel_segment_between(prev.id, entry.id)
                            if seg:
                                route_items.append(("travel", seg))
                            incoming_mode_label = _pdf_mode_label(seg.mode if seg else None)
                        else:
                            seg = route.travel_segment_between(prev.id, entry.id)
                            incoming_mode_label = _pdf_mode_label(seg.mode if seg else None)
                        if self._settings.show_space_blocks:
                            esp = route.empty_space_between(prev.id, entry.id)
                            if esp and esp.duration_minutes > 0:
                                route_items.append(("empty", esp))
                    route_items.append(("visit", (entry, incoming_mode_label)))

                if not route_items:
                    route_chunk_map[route.id] = [[]]
                    continue

                current_chunk: list[tuple[str, object]] = []
                current_h = 0
                route_chunks: list[list[tuple[str, object]]] = []
                for item_type, item in route_items:
                    item_h = visit_h
                    if item_type == "travel":
                        item_h = travel_h
                    elif item_type == "empty":
                        item_h = empty_h
                    elif item_type == "extra":
                        item_h = extra_h

                    block_h = item_h + block_gap
                    if current_chunk and current_h + block_h > available_h:
                        route_chunks.append(current_chunk)
                        current_chunk = []
                        current_h = 0
                    current_chunk.append((item_type, item))
                    current_h += block_h

                if current_chunk:
                    route_chunks.append(current_chunk)

                route_chunk_map[route.id] = route_chunks

            first_route = True
            for route in routes:
                chunks = route_chunk_map.get(route.id, [[]])

                if not first_route:
                    writer.newPage()
                first_route = False

                page_local_chunk_idx = 0
                while page_local_chunk_idx < len(chunks):
                    if page_local_chunk_idx > 0:
                        writer.newPage()

                    col_x_left = left
                    draw_route_column(
                        route,
                        col_x_left,
                        top,
                        chunks[page_local_chunk_idx],
                        continued=(page_local_chunk_idx > 0),
                    )
                    page_local_chunk_idx += 1

                    if page_local_chunk_idx < len(chunks):
                        col_x_right = left + (column_w + column_gap)
                        draw_route_column(
                            route,
                            col_x_right,
                            top,
                            chunks[page_local_chunk_idx],
                            continued=(page_local_chunk_idx > 0),
                        )
                        page_local_chunk_idx += 1
        finally:
            painter.end()

    # ------------------------------------------------------------------
    # Auto-save
    # ------------------------------------------------------------------

    @Slot()
    def _on_autosave(self):
        self._db.save_settings(self._settings)
        dirty = self._autosave.dirty_routes()
        for route_id in dirty:
            route = self._routes.get(route_id)
            if route:
                self._db.save_route(route)
        self._autosave.clear()

    # ------------------------------------------------------------------
    # Travel status window
    # ------------------------------------------------------------------

    def _get_travel_status(self) -> TravelStatusDialog:
        if self._travel_status is None:
            self._travel_status = TravelStatusDialog(self)
            self._travel_status.configure_file_logging(
                self._settings.file_logging_enabled,
                self._settings.file_logging_retention_days,
            )
        return self._travel_status

    @Slot()
    def _show_travel_status(self):
        dlg = self._get_travel_status()
        dlg.show()
        dlg.raise_()

    # ------------------------------------------------------------------
    # Visit selection + pair highlighting
    # ------------------------------------------------------------------

    @Slot(int, int)
    def _on_route_visit_selected(self, route_id: int, entry_id: int):
        self._route_scene.set_selected_entry(route_id, entry_id)
        self._pool_scene.set_selected_visit(None)
        route = self._routes.get(route_id)
        if route:
            entry = next((e for e in route.entries if e.id == entry_id), None)
            if entry and entry.visit_id:
                self._selected_visit_id = entry.visit_id
                self._sync_map_selection()
                self._highlight_pair_for_visit(entry.visit_id)
                return
        self._selected_visit_id = None
        self._sync_map_selection()
        self._highlight_pair_for_visit(None)

    @Slot(int)
    def _on_pool_visit_selected(self, visit_id: int):
        self._pool_scene.set_selected_visit(visit_id)
        self._route_scene.set_selected_entry(None, None)
        self._selected_visit_id = visit_id
        self._sync_map_selection()
        self._highlight_pair_for_visit(visit_id)

    def _highlight_pair_for_visit(self, visit_id: Optional[int]):
        partner_id = self._pairs.get(visit_id) if visit_id else None
        self._route_scene.highlight_pair(None, None)
        self._pool_scene.highlight_pair(partner_id)
        if partner_id:
            # Find partner in routes
            for route in self._routes.values():
                for entry in route.entries:
                    if entry.visit_id == partner_id:
                        self._route_scene.highlight_pair(route.id, entry.id)

    @Slot(int, int)
    def _on_route_entry_pair_requested(self, route_id: int, entry_id: int):
        """Handle pair request from route visit."""
        route = self._routes.get(route_id)
        if not route:
            return
        
        entry = next((e for e in route.entries if e.id == entry_id), None)
        if not entry or not entry.visit_id:
            return
        
        self._show_pair_dialog(entry.visit_id)

    @Slot(int, int)
    def _on_route_entry_unpair_requested(self, route_id: int, entry_id: int):
        """Handle unpair request from route visit."""
        route = self._routes.get(route_id)
        if not route:
            return
        
        entry = next((e for e in route.entries if e.id == entry_id), None)
        if not entry or not entry.visit_id:
            return
        
        self._unpair_visit(entry.visit_id)

    @Slot(int)
    def _on_pool_visit_pair_requested(self, visit_id: int):
        """Handle pair request from pool visit."""
        self._show_pair_dialog(visit_id)

    @Slot(int)
    def _on_pool_visit_unpair_requested(self, visit_id: int):
        """Handle unpair request from pool visit."""
        self._unpair_visit(visit_id)

    @Slot(int)
    def _on_template_edit_requested(self, template_index: int):
        if template_index < 0 or template_index >= len(self._default_templates):
            return

        template = self._default_templates[template_index]
        if template.get("type") == "extra_time":
            return

        dlg = QDialog(self)
        dlg.setWindowTitle("Redigera standardbesök")
        form = QFormLayout(dlg)

        name_edit = QLineEdit(str(template.get("name", "Kontor")), dlg)
        address_edit = QLineEdit(str(template.get("address", "")), dlg)
        duration_spin = QSpinBox(dlg)
        duration_spin.setRange(1, 720)
        duration_spin.setValue(max(1, int(template.get("duration_minutes", 10))))
        duration_spin.setSuffix(" min")

        form.addRow("Namn:", name_edit)
        form.addRow("Adress:", address_edit)
        form.addRow("Duration:", duration_spin)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel,
            parent=dlg,
        )
        form.addRow(buttons)
        buttons.accepted.connect(dlg.accept)
        buttons.rejected.connect(dlg.reject)

        if dlg.exec() != QDialog.DialogCode.Accepted:
            return

        name_value = (name_edit.text() or "").strip() or "Kontor"
        address_value = (address_edit.text() or "").strip()
        if not address_value:
            QMessageBox.warning(self, "Ogiltig adress", "Adress kan inte vara tom.")
            return

        template["name"] = name_value
        template["address"] = address_value
        template["duration_minutes"] = int(duration_spin.value())

        placed: set[int] = set()
        for route in self._routes.values():
            for entry in route.entries:
                if entry.visit_id:
                    placed.add(entry.visit_id)
        pool_visits = [v for v in self._visits.values() if v.id not in placed]
        pool_order = {street: order for street, order in self._pool_scene.column_order_pairs()}
        self._pool_scene.load(self._default_templates, pool_visits)
        self._pool_scene.set_column_order_map(pool_order)
        self._pool_scene.set_extra_time_minutes(self._settings.extra_time_minutes)
        self._apply_filters_and_search()
        self._sync_map_window_visits()

        if self._api_key and address_value:
            try:
                MapGeocodingService(self._db, self._api_key).precache_addresses([address_value])
            except Exception:
                pass

    def _show_pair_dialog(self, visit_id: int):
        """Show dialog to pair visit with another visit."""
        visit = self._visits.get(visit_id)
        if not visit:
            return

        sel_name = (visit.name or "").strip().lower()
        sel_address = (visit.address or "").strip().lower()
        available = [
            v for v in self._visits.values()
            if (v.id != visit_id and v.id not in self._pairs and
                (v.name or "").strip().lower() == sel_name and
                (v.address or "").strip().lower() == sel_address)
        ]
        source_route, _ = self._find_route_entry_by_visit(visit_id)
        route_by_visit: dict[int, Optional[int]] = {}
        current_interval_by_visit: dict[int, Optional[str]] = {}
        recommended_visit_id: Optional[int] = None
        for candidate in available:
            candidate_route, _ = self._find_route_entry_by_visit(candidate.id)
            route_by_visit[candidate.id] = candidate_route.id if candidate_route else None
            _, candidate_entry = self._find_route_entry_by_visit(candidate.id)
            if candidate_entry:
                current_interval_by_visit[candidate.id] = f"{candidate_entry.start_time}–{candidate_entry.end_time}"
            if recommended_visit_id is None and self._is_recommended_pair_candidate(visit, candidate):
                recommended_visit_id = candidate.id

        _, selected_entry = self._find_route_entry_by_visit(visit_id)
        if selected_entry:
            current_interval_by_visit[visit_id] = f"{selected_entry.start_time}–{selected_entry.end_time}"
        
        from ui.dialogs.pair_dialog import PairDialog
        dialog = PairDialog(
            available,
            visit,
            source_route_id=source_route.id if source_route else None,
            route_by_visit=route_by_visit,
            current_interval_by_visit=current_interval_by_visit,
            recommended_visit_id=recommended_visit_id,
            parent=self,
        )
        if dialog.exec():
            selected_id = dialog.get_selected_visit_id()
            if selected_id:
                self._pair_visits(visit_id, selected_id)

    def _rebuild_all_views(self):
        pool_order = {street: order for street, order in self._pool_scene.column_order_pairs()}
        placed: set[int] = set()
        for route in self._routes.values():
            for entry in route.entries:
                if entry.visit_id:
                    placed.add(entry.visit_id)

        pool_visits = [v for v in self._visits.values() if v.id not in placed]
        self._pool_scene.load(self._default_templates, pool_visits)
        self._pool_scene.set_column_order_map(pool_order)
        self._route_scene.load_routes(list(self._routes.values()))
        self._pool_scene.set_extra_time_minutes(self._settings.extra_time_minutes)
        self._route_scene.set_extra_time_minutes(self._settings.extra_time_minutes)
        self._on_block_visibility_changed()
        self._route_scene.set_pair_alignment_enabled(self._settings.align_pair_visits)
        self._apply_filters_and_search()
        self._sync_map_window_visits()
        self._sync_map_selection()
        self._highlight_pair_for_visit(self._selected_visit_id)

    def _pair_visits(self, visit_id_a: int, visit_id_b: int):
        """Create a pairing between two visits."""
        route_a, _ = self._find_route_entry_by_visit(visit_id_a)
        route_b, _ = self._find_route_entry_by_visit(visit_id_b)
        if route_a and route_b and route_a.id == route_b.id:
            QMessageBox.warning(
                self,
                "Kan inte para",
                "Besök på samma rutt kan inte paras.",
            )
            return
        self._pairs[visit_id_a] = visit_id_b
        self._pairs[visit_id_b] = visit_id_a
        self._sync_visit_pair_state()
        self._rebuild_all_views()

    def _unpair_visit(self, visit_id: int):
        """Remove pairing for a visit."""
        partner_id = self._pairs.get(visit_id)
        if partner_id:
            del self._pairs[visit_id]
            del self._pairs[partner_id]
            self._sync_visit_pair_state()
            self._rebuild_all_views()

    def _highlight_pair_for_visit(self, visit_id: Optional[int]):
        partner_id = self._pairs.get(visit_id) if visit_id else None
        self._route_scene.highlight_pair(None, None)
        self._pool_scene.highlight_pair(partner_id)
        if partner_id:
            # Find partner in routes
            for route in self._routes.values():
                for entry in route.entries:
                    if entry.visit_id == partner_id:
                        self._route_scene.highlight_pair(route.id, entry.id)
        
        # Highlight same-name-address visits
        self._pool_scene.highlight_same_name_address(visit_id, self._visits)
        self._route_scene.highlight_same_name_address(visit_id, self._visits)

    def _find_route_entry_by_visit(self, visit_id: int) -> tuple[Optional[Route], Optional[RouteEntry]]:
        for route in self._routes.values():
            for entry in route.entries:
                if entry.visit_id == visit_id:
                    return route, entry
        return None, None

    def _route_contains_visit(self, route: Route, visit_id: int) -> bool:
        return any(entry.visit_id == visit_id for entry in route.entries)

    def _can_place_paired_visit_on_route(self, route: Route, visit_id: int) -> bool:
        partner_id = self._pairs.get(visit_id)
        if not partner_id:
            return True
        return not self._route_contains_visit(route, partner_id)

    def _pair_earliest_start(self, route: Route, entry: RouteEntry) -> int:
        entries = route.sorted_entries()
        idx = next((i for i, e in enumerate(entries) if e.id == entry.id), None)
        if idx is None or idx == 0:
            return 0
        prev = entries[idx - 1]
        seg = route.travel_segment_between(prev.id, entry.id)
        travel = max(0, seg.travel_minutes) if seg else 0
        extra = self._settings.extra_time_minutes if route.extra_time_for_entry(entry.id) else 0
        return _t2m(prev.end_time) + travel + extra

    def _set_entry_start_with_gap(self, route: Route, entry: RouteEntry, target_start: int) -> bool:
        entries = route.sorted_entries()
        idx = next((i for i, e in enumerate(entries) if e.id == entry.id), None)
        if idx is None:
            return False

        duration = max(1, _t2m(entry.end_time) - _t2m(entry.start_time))
        current_start = _t2m(entry.start_time)

        if idx == 0:
            new_start = max(0, int(target_start))
            if new_start == current_start:
                return False
            entry.start_time = _m2t(new_start)
            entry.end_time = _m2t(new_start + duration)
            self._recalc.recalculate(route)
            return True

        prev = entries[idx - 1]
        base_start = self._pair_earliest_start(route, entry)
        new_start = max(base_start, int(target_start))
        gap = max(0, new_start - base_start)

        space = route.empty_space_between(prev.id, entry.id)
        old_gap = max(0, space.duration_minutes) if space else 0
        if new_start == current_start and old_gap == gap:
            return False

        if gap > 0:
            if space is None:
                route.empty_spaces.append(EmptySpace(
                    id=None,
                    route_id=route.id,
                    from_entry_id=prev.id,
                    to_entry_id=entry.id,
                    duration_minutes=gap,
                ))
            else:
                space.duration_minutes = gap
        elif space is not None:
            route.empty_spaces.remove(space)

        entry.start_time = _m2t(new_start)
        entry.end_time = _m2t(new_start + duration)
        self._recalc.recalculate(route)
        return True

    def _synchronize_pair_for_visit(self, visit_id: int) -> set[int]:
        partner_id = self._pairs.get(visit_id)
        if not partner_id:
            return set()

        route_a, entry_a = self._find_route_entry_by_visit(visit_id)
        route_b, entry_b = self._find_route_entry_by_visit(partner_id)
        if not route_a or not entry_a or not route_b or not entry_b:
            return set()

        changed_routes: set[int] = set()

        earliest_a = self._pair_earliest_start(route_a, entry_a)
        earliest_b = self._pair_earliest_start(route_b, entry_b)
        current_a = _t2m(entry_a.start_time)
        current_b = _t2m(entry_b.start_time)
        target_start = max(earliest_a, earliest_b, current_a, current_b)

        if self._set_entry_start_with_gap(route_a, entry_a, target_start):
            changed_routes.add(route_a.id)
        if self._set_entry_start_with_gap(route_b, entry_b, target_start):
            changed_routes.add(route_b.id)

        return changed_routes

    def _align_pair_to_start(self, visit_id: int, target_start: int) -> set[int]:
        partner_id = self._pairs.get(visit_id)
        if not partner_id:
            return set()

        route_a, entry_a = self._find_route_entry_by_visit(visit_id)
        route_b, entry_b = self._find_route_entry_by_visit(partner_id)
        if not route_a or not entry_a or not route_b or not entry_b:
            return set()

        changed_routes: set[int] = set()
        resolved_target = max(
            int(target_start),
            self._pair_earliest_start(route_a, entry_a),
            self._pair_earliest_start(route_b, entry_b),
        )

        if self._set_entry_start_with_gap(route_a, entry_a, resolved_target):
            changed_routes.add(route_a.id)
        if self._set_entry_start_with_gap(route_b, entry_b, resolved_target):
            changed_routes.add(route_b.id)

        return changed_routes

    def _pair_start_snapshot(self) -> tuple[tuple[int, int, int, int], ...]:
        rows: list[tuple[int, int, int, int]] = []
        for visit_id, partner_id in self._pairs.items():
            if visit_id >= partner_id:
                continue
            route_a, entry_a = self._find_route_entry_by_visit(visit_id)
            route_b, entry_b = self._find_route_entry_by_visit(partner_id)
            if not route_a or not entry_a or not route_b or not entry_b:
                continue
            rows.append((visit_id, partner_id, _t2m(entry_a.start_time), _t2m(entry_b.start_time)))
        return tuple(sorted(rows))

    def _synchronize_all_pairs(self, max_passes: int = 6) -> set[int]:
        changed_routes: set[int] = set()
        pair_roots = [visit_id for visit_id, partner_id in self._pairs.items() if visit_id < partner_id]
        prev_snapshot = self._pair_start_snapshot()
        for _ in range(max(1, int(max_passes))):
            changed_this_pass: set[int] = set()
            for visit_id in pair_roots:
                changed_this_pass.update(self._synchronize_pair_for_visit(visit_id))
            changed_routes.update(changed_this_pass)
            current_snapshot = self._pair_start_snapshot()
            if not changed_this_pass or current_snapshot == prev_snapshot:
                break
            prev_snapshot = current_snapshot
        return changed_routes

    def _apply_pair_synchronization_and_refresh(self, max_passes: int = 6):
        changed_route_ids = self._synchronize_all_pairs(max_passes=max_passes)
        for changed_route_id in changed_route_ids:
            changed_route = self._routes.get(changed_route_id)
            if changed_route:
                self._route_scene.rebuild_route(changed_route, animate=False)
                self._autosave.mark_dirty(changed_route_id)

    # ------------------------------------------------------------------
    # Close event
    # ------------------------------------------------------------------

    def closeEvent(self, event):
        if self._map_window is not None:
            self._map_window.close()
        self._autosave.flush_now()
        self._db.save_settings(self._settings)
        self._db.close()
        super().closeEvent(event)
